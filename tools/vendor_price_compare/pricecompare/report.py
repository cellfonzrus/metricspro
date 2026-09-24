"""Write the comparison workbook (Excel) + CSV copies."""
from __future__ import annotations

import csv
from pathlib import Path

from . import core

MONEY = '"$"#,##0.00'
UNIT_MONEY = '"$"#,##0.0000'
PCT = "0.0%"


def _autosize(ws, widths=None):
    from openpyxl.utils import get_column_letter
    for i, col in enumerate(ws.iter_cols(min_row=1, max_row=min(ws.max_row, 300)), start=1):
        w = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, w + 2), (widths or {}).get(i, 60))


def _header(ws, cols):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="351C15")   # brown header
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "A2"


def write_workbook(path, vendors, products, compared, singles, plan, run_info):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    green = PatternFill("solid", fgColor="C6EFCE")
    amber = PatternFill("solid", fgColor="FFEB9C")
    names = {v["key"]: v.get("name", v["key"]) for v in vendors}
    keys = [v["key"] for v in vendors]

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Vendor price comparison"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append(["Run at", run_info["run_at"]])
    ws.append([])
    ws.append(["Vendor", "Logged in", "Pages read", "Products found", "Notes"])
    for c in ws[4]:
        c.font = Font(bold=True)
    for v in vendors:
        r = run_info["vendors"].get(v["key"], {})
        ws.append([names[v["key"]], "yes" if r.get("login_ok") else "NO", r.get("pages", 0), r.get("products", 0),
                   " | ".join(r.get("notes") or [])])
    ws.append([])
    ws.append(["Products sold by 2+ vendors (compared)", len(compared)])
    ws.append(["Products found at one vendor only", len(singles)])
    tot = sum(r["savings"] or 0 for r in compared if r["basis"] == "listed price")
    ws.append(["Sum of per-item price gaps (listed price, 1 of each)", tot])
    ws.cell(ws.max_row, 2).number_format = MONEY
    ws.append([])
    ws.append(["How to read this"])
    ws.cell(ws.max_row, 1).font = Font(bold=True)
    for line in [
        "Comparison: one row per product found at more than one vendor. Green = cheapest in-stock vendor.",
        "Match column: 'exact' = same item/part number; 'likely' = same size + very similar name; "
        "'possible — check' = please eyeball before ordering.",
        "Basis: 'per unit' when every vendor states the pack size (e.g. 25/bundle), else the listed price.",
        "Availability 'unknown' means the vendor page did not say — it is NOT assumed in stock.",
        "Order plan: only filled when shopping_list.csv has rows.",
        "Prices are what your logged-in account saw at the run time above. Confirm in the cart before paying.",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["E"].width = 80

    # Comparison
    wc = wb.create_sheet("Comparison")
    cols = ["Product", "Match", "Match why", "Basis", "Cheapest vendor", "Best price", "Next price", "Save $", "Save %"]
    for k in keys:
        cols += [f"{names[k]} item #", f"{names[k]} price", f"{names[k]} pack", f"{names[k]} stock", f"{names[k]} link"]
    cols.append("Note")
    _header(wc, cols)
    for r in compared:
        row = [r["product"], r["match"], r["match_method"], r["basis"], names.get(r["best_vendor"], r["best_vendor"]),
               r["best_price"], r["next_price"], r["savings"], r["savings_pct"]]
        for k in keys:
            p = r.get(k)
            row += ([p.get("sku"), p.get("price"), p.get("pack_qty"), p.get("availability"), p.get("url")] if p
                    else [None, None, None, "not sold here", None])
        row.append(r["note"])
        wc.append(row)
        rn = wc.max_row
        fmt = UNIT_MONEY if r["basis"] == "per unit" else MONEY
        for ci in (6, 7, 8):
            wc.cell(rn, ci).number_format = fmt
        wc.cell(rn, 9).number_format = PCT
        for i, k in enumerate(keys):
            pc = wc.cell(rn, 10 + i * 5 + 1)
            pc.number_format = MONEY
            if k == r["best_vendor"]:
                pc.fill = green
        if r["match"].startswith("possible"):
            wc.cell(rn, 2).fill = amber
    wc.auto_filter.ref = wc.dimensions
    _autosize(wc, {1: 55})

    # All products
    wa = wb.create_sheet("All products")
    acol = ["Vendor", "Product", "Item #", "Price", "Was (list) price", "Pack qty", "Unit price", "Availability",
            "Stock qty", "Stock text", "Link"]
    _header(wa, acol)
    for p in sorted(products, key=lambda p: (p["vendor"], p["name"].lower())):
        wa.append([names.get(p["vendor"], p["vendor"]), p["name"], p["sku"], p["price"], p["list_price"], p["pack_qty"],
                   core.unit_price(p), p["availability"], p["stock_qty"], p["stock_text"], p["url"]])
        rn = wa.max_row
        wa.cell(rn, 4).number_format = MONEY
        wa.cell(rn, 5).number_format = MONEY
        wa.cell(rn, 7).number_format = UNIT_MONEY
    wa.auto_filter.ref = wa.dimensions
    _autosize(wa, {2: 60})

    # One vendor only
    wo = wb.create_sheet("One vendor only")
    _header(wo, ["Vendor", "Product", "Item #", "Price", "Availability", "Link"])
    for p in sorted(singles, key=lambda p: (p["vendor"], p["name"].lower())):
        wo.append([names.get(p["vendor"], p["vendor"]), p["name"], p["sku"], p["price"], p["availability"], p["url"]])
        wo.cell(wo.max_row, 4).number_format = MONEY
    wo.auto_filter.ref = wo.dimensions
    _autosize(wo, {2: 60})

    # Order plan
    wp = wb.create_sheet("Order plan")
    pcol = ["You searched", "Qty (as the vendor sells it)", "Pick vendor", "Pick product", "Pick item #", "Pack", "Price", "Line total", "Stock"]
    for k in keys:
        pcol += [f"{names[k]} match", f"{names[k]} price"]
    pcol.append("Note")
    _header(wp, pcol)
    totals = {}
    for line in plan:
        pk = line["pick"]
        row = [line["search"], line["qty"], names.get(pk["vendor"]) if pk else None, pk["name"] if pk else None,
               pk["sku"] if pk else None, pk["pack_qty"] if pk else None, pk["price"] if pk else None,
               line["line_total"], pk["availability"] if pk else None]
        byv = {p["vendor"]: p for p in line["options"]}
        for k in keys:
            p = byv.get(k)
            row += [p["name"] if p else None, p["price"] if p else None]
        row.append(line["note"])
        wp.append(row)
        rn = wp.max_row
        wp.cell(rn, 7).number_format = MONEY
        wp.cell(rn, 8).number_format = MONEY
        for i in range(len(keys)):
            wp.cell(rn, 11 + i * 2).number_format = MONEY
        if pk:
            totals[pk["vendor"]] = totals.get(pk["vendor"], 0) + (line["line_total"] or 0)
    if plan:
        wp.append([])
        for k, t in totals.items():
            wp.append([f"Order total at {names.get(k, k)}", None, None, None, None, None, None, round(t, 2)])
            wp.cell(wp.max_row, 8).number_format = MONEY
            wp.cell(wp.max_row, 1).font = Font(bold=True)
    else:
        wp.append(["(empty — put items in shopping_list.csv and run again with --compare-only to fill this sheet)"])
    _autosize(wp, {1: 40, 4: 50})

    wb.save(path)


def write_csv(path, rows, cols):
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=cols, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)
