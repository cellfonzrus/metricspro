"""DB-FREE PROOF — the tenant-onboarding intake, STAGE D-2: step 2.5a "What counts as an activation"
(owner 2026-09-21: "sales report shows 88 txns but not a break up in to activations and upgrade etc,
also nothing on exec mtd").

    python3 harness_onboarding_intake_d.py        (from backend/; no network, no DB)

THE SCENARIO, over the in-memory supabase client (harness_intake_fakes) driving the REAL router:
    §A analyze a line-level sales export whose contract_type column does not exist and whose activation
       TYPE sits in the category path leaf → the payload carries `line_class`: current rules classify
       nothing (gate OPEN), the proposal per class with the count it would classify, the metric buckets
    §B commit → the rows LAND (ok), but the export is NOT verified: needs_input, the blocking reason
       names step 2.5a; the Stage-4 row is red, fix_step 2.5a, with the activation note
    §C GET /onboarding/intake/line-class over the LANDED slice (re-read as the commit re-read it)
    §D PUT /onboarding/intake/line-class with the proposal (+ the phones metric rule) → written into the
       TWO existing homes only (accessory_config.activation_details_rules — the Activation-Details basis
       keys beside it preserved; exec_metric_config 'phones'), the slice RE-COUNTED, the row verified
    §E every reader dereferences the saved rules: _accessory_config['line_rules'], _sales_cell_agg
       (Sales Report / Exec MTD / Daily Targets) counts 39 / 25 / 23 by distinct invoice and 67 phones,
       the Executive MTD endpoint's `landing.classified` gate is closed, the pay path's cfg carries them
    §F the attestation path: an export with genuinely no activations is verified only with a reason
       and a name; the Exec-MTD `landing.classified` banner names the step for an unmapped tenant
    §G negative controls: junk bodies refused; an unknown metric bucket refused; a non-existent
       instance 404; the house org's own rows byte-identical (no config → house defaults)
"""
import asyncio
import io
import json
import os
import sys
from datetime import date

sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from harness_intake_fakes import FakeDB, FakeUpload, measured_fixture_rows, CAT   # noqa: E402
from app.modules.commcalc import onboarding_intake as OI                            # noqa: E402
from app.modules.commcalc import line_class as LC                                   # noqa: E402
from app.modules.commcalc import router as R                                        # noqa: E402
from app.modules.storeops import router as SO                                       # noqa: E402
from app.modules.storeops import merchant_ids as MIDS                               # noqa: E402
from app.modules.closing import router as CR                                        # noqa: E402
from app.modules.commcalc import epay_ingest as EPI                                 # noqa: E402

ORG = "d2d2d2d2-0000-4000-8000-00000000d2d2"
HOUSE = "00000000-0000-0000-0000-000000000001"
_pass = _fail = 0


def check(name, cond, extra=""):
    global _pass, _fail
    if cond:
        _pass += 1
        print(f"  PASS  {name}")
    else:
        _fail += 1
        print(f"  FAIL  {name}{(' — ' + str(extra)[:700]) if extra else ''}")
    return bool(cond)


def section(t):
    print(f"\n{t}\n" + "─" * len(t))


def use(db):
    R.sb = lambda: db
    SO.sb = lambda: db
    SO.get_supabase = lambda: db
    MIDS.get_supabase = lambda: db
    CR.get_supabase = lambda: db
    EPI.get_supabase = lambda: db
    R._TABLE_COL_PRESENT.clear()
    R._invalidate_accessory_config(ORG)
    return db


def run(coro):
    return asyncio.run(coro)


def analyze(db, data, filename, **form):
    use(db)
    kw = dict(org_id=ORG)
    kw.update(form)
    return run(R.onboarding_intake_analyze(file=FakeUpload(data, filename), **kw))


def commit(db, data, filename, **form):
    use(db)
    kw = dict(verified_by="tester", org_id=ORG)
    kw.update(form)
    return run(R.onboarding_intake_commit(file=FakeUpload(data, filename), **kw))


def http_error(fn, *a, **kw):
    try:
        fn(*a, **kw)
        return None, None
    except R.HTTPException as e:
        return e.status_code, e.detail


def fresh_db():
    db = FakeDB()
    db.seed("stores", [{"org_id": ORG, "store_code": "V-10", "address": "10 Main St", "market": "NY", "is_active": True},
                       {"org_id": ORG, "store_code": "V-22", "address": "22 Oak Ave", "market": "NY", "is_active": True}])
    db.seed("store_mapping", [{"org_id": ORG, "store_code": "V-10", "store_address": "10 Main St", "market": "NY"},
                              {"org_id": ORG, "store_code": "V-22", "store_address": "22 Oak Ave", "market": "NY"}])
    db.seed("employees", [{"org_id": ORG, "employee_id": "E1", "name": "alice", "epay_salesperson": "alice", "is_active": True},
                          {"org_id": ORG, "employee_id": "E2", "name": "bob", "epay_salesperson": "bob", "is_active": True},
                          {"org_id": ORG, "employee_id": "E3", "name": "carol", "epay_salesperson": "carol", "is_active": True}])
    db.seed("carrier", [{"id": "aaaaaaaa-0000-0000-0000-00000000c0d2", "org_id": ORG, "name": "Northwind Cellular", "code": "northwind"}])
    # the org already carries an Activation-Details basis rule in the SAME JSON column — it must survive the 2.5a save
    db.seed("accessory_config", [{"org_id": ORG, "departments": [], "categories": [], "product_keywords": [], "acima_tenders": [],
                                  "activation_details_rules": {"edge_name_tokens": ["edge plan"]}}])
    return db


# the measured export: the headers as the saved mapping names them, the store under a known synonym
HEADERS = ["Sold On", "Invoiced At", "Tendered By", "Sold By", "Invoice #", "Customer", "Product SKU", "Product Name",
           "Category", "Quantity", "Total Price", "Total Cost", "Gross Profit", "Tracking #", "Refund"]
ROWS = measured_fixture_rows()
FILE_TOTAL = round(sum(r["ext_price"] for r in ROWS), 2)


def sales_xlsx(rows=ROWS):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(HEADERS)
    for i, r in enumerate(rows):
        ws.append([r["trans_date"], r["store"], r["salesperson"], r["salesperson"], r["trans_id"], "Walk-in", f"SKU{i}",
                   r["product_desc"], r["category"], 1, f"{r['ext_price']:.2f}", f"{r['ext_price'] * 0.7:.2f}",
                   f"{r['gp']:.2f}", r["serial_1"], "No"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SALES_MAP = {"trans_id": "Invoice #", "store": "Invoiced At", "salesperson": "Sold By", "user_login": "Tendered By",
             "trans_date": "Sold On", "sku": "Product SKU", "serial_1": "Tracking #", "product_desc": "Product Name",
             "category": "Category", "department": "Category", "ext_price": "Total Price", "gp": "Gross Profit",
             "voided": "Refund", "quantity": "Quantity", "total_cost": "Total Cost"}
IKEY = OI.stage2_instance_key("sales", "mypos", "sales")

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§A analyze — the payload carries the 2.5a block")
db = fresh_db()
a = analyze(db, sales_xlsx(), "sales.xlsx", source_kind="sales", pos_source="mypos", layout="sales",
            column_map=json.dumps(SALES_MAP), typed_total=f"{FILE_TOTAL:.2f}")
lc = a.get("line_class") or {}
check("A1 analyze proposes the columns (contract_type unmapped — the file has none) and carries `line_class`",
      not any(c["target_field"] == "contract_type" and c.get("column") for c in a["columns"]) and lc.get("step") == "2.5a", (lc.get("error"), [c["target_field"] for c in a["columns"] if c.get("column")]))
check("A2 under the org's CURRENT (house) rules nothing classifies: gate OPEN, the note names the fields read and the step",
      lc["gate_open"] is True and lc["current"]["activation_type_lines"] == 0 and lc["current"]["scanned"] == len(ROWS)
      and "contract_type" in lc["gate_note"] and "2.5a" in lc["gate_note"], (lc.get("current"), lc.get("gate_note")))
prop = lc["suggest"]["proposal"]
check("A3 the proposal names category / product words per class with their counts; the preview is 40 / 25 / 24 / 4 lines",
      "category" in prop["fields"] and "new activation" in prop["tokens"]["activation"] and "upgrade" in prop["tokens"]["upgrade"]
      and {"customer provided", "customer owned"} <= set(prop["tokens"]["byod"]) and prop["tokens"]["hardware_only"]
      and lc["suggest"]["preview"]["classes"]["activation"]["lines"] == 40 and lc["suggest"]["preview"]["classes"]["upgrade"]["lines"] == 25
      and lc["suggest"]["preview"]["classes"]["byod"]["lines"] == 24 and lc["suggest"]["preview"]["classes"]["hardware_only"]["lines"] == 4, prop)
mb = (lc.get("metrics") or {}).get("buckets") or {}
check("A4 the metric buckets: phones matched 0 and proposes category_contains words → 67; bill_payment matched 0 and proposes nothing",
      mb.get("phones", {}).get("matched") == 0 and mb["phones"]["preview"] == 67 and mb["bill_payment"]["matched"] == 0 and mb["bill_payment"]["proposal"] is None, mb.get("phones"))
check("A5 the step exists on the spine after 2.5 (rail vocabulary + next_step); since 2026-09-21 the spine's next key is 2.5b (the invoice export's tender columns) — the page walks a SALES export from 2.5a straight to 2.6 by kind, as it walks an invoice export past 2.5a",
      OI.STEP_KEYS_STAGE2[OI.STEP_KEYS_STAGE2.index("2.5") + 1] == "2.5a" and OI.next_step("2.5") == "2.5a" and OI.next_step("2.5a") == "2.5b"
      and OI.next_step("2.5b") == "2.6" and "2.5a" in OI.ALL_STEP_KEYS)

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§B commit — landed, but NOT verified until the words are mapped")
c = commit(db, sales_xlsx(), "sales.xlsx", source_kind="sales", pos_source="mypos", layout="sales",
           column_map=json.dumps(SALES_MAP), typed_total=f"{FILE_TOTAL:.2f}")
vn = c["verified_numbers"]
check("B1 the landing is ok (rows re-read = rows built, totals tie) — the rows ARE in the table",
      c["ok"] is True and vn["rows_landed"] == len(ROWS) and vn["tie"]["match"] is True, (c.get("problems"), vn.get("rows_landed")))
check("B2 …but `verified` is False: the gate is blocked, the problem names step 2.5a, the stage row is needs_input with that reason",
      c["verified"] is False and c["activation_gate"]["blocked"] is True and any("2.5a" in p for p in c["problems"])
      and vn["activation_classes"]["activation_type_lines"] == 0 and vn["landing_ok"] is True, (c.get("problems"), c.get("activation_gate")))
st = R.onboarding_intake_state(org_id=ORG)
row = next(r for r in st["rail"]["verify_table"] if r["instance_key"] == IKEY)
check("B3 the Stage-4 row is RED, fix_step 2.5a, blocking reason on the row, the activation note says no line could be told apart",
      row["red"] is True and row["fix_step"] == "2.5a" and "2.5a" in (row["blocking_reason"] or "")
      and "no line could be told apart" in (row["activation_note"] or ""), row)
inst = next(i for i in st["rail"]["instances"] if i["instance_key"] == IKEY)
check("B4 the instance's status is needs_input (not verified); the stage row carries the counts", inst["status"] == "needs_input")

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§C GET /onboarding/intake/line-class — over the LANDED slice")
g = R.onboarding_intake_line_class(instance_key=IKEY, org_id=ORG)
check("C1 landed:true, the re-read slice is the whole file, the block's current count is 0 / gate open, the recorded count matches",
      g["landed"] is True and g["rows"] == len(ROWS) and g["block"]["gate_open"] is True
      and g["recorded"]["activation_type_lines"] == 0 and g["status"] == "needs_input", (g.get("rows"), g.get("status")))
check("C2 the proposal over the landed rows is the same proposal analyze made (same rows, same engine)",
      g["block"]["suggest"]["proposal"] == prop)
status, detail = http_error(R.onboarding_intake_line_class, instance_key="sales:nope:sales", org_id=ORG)
check("C3 an unknown instance → 404", status == 404, (status, detail))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§D PUT /onboarding/intake/line-class — saved into the two homes, re-counted, verified")
R._can_edit_classification = lambda *_a, **_k: True      # the classification permission (the gate is core's; the harness grants it)
body = R.OnboardingLineClassIn(instance_key=IKEY, fields=prop["fields"], tokens=prop["tokens"],
                               metric_rules={"phones": mb["phones"]["proposal"]}, by="tester")
p = R.onboarding_intake_put_line_class(body, org_id=ORG)
check("D1 the save answers ok: activation rules written, phones written, the slice re-counted, verified:true, gate closed",
      p["ok"] and p["written"]["activation_rules"] is True and p["written"]["metric_buckets"] == ["phones"]
      and p["landed"] is True and p["verified"] is True and p["activation_gate"]["blocked"] is False
      and p["block"]["current"]["classes"]["activation"]["transactions"] == 39, (p.get("written"), p.get("activation_gate")))
acc_row = next(r for r in db.tables["accessory_config"] if r["org_id"] == ORG)
adr = acc_row["activation_details_rules"]
check("D2 accessory_config.activation_details_rules holds fields + tokens (ours) AND the Activation-Details basis key that was there before",
      adr["fields"] == prop["fields"] and adr["tokens"]["activation"] == prop["tokens"]["activation"]
      and adr["tokens"]["byod"] == prop["tokens"]["byod"] and adr["edge_name_tokens"] == ["edge plan"], adr)
emc = [r for r in db.tables.get("exec_metric_config", []) if r["org_id"] == ORG]
check("D3 exec_metric_config carries ONE row for phones with the contains tokens; no other bucket was written",
      len(emc) == 1 and emc[0]["bucket"] == "phones" and {"smartphone", "basic phone"} <= set(emc[0]["rules"].get("category_contains") or []), emc)
check("D4 no other table gained a row from the save (the stage row was UPDATED, not inserted twice)",
      len([r for r in db.tables["onboarding_stage_state"] if r["instance_key"] == IKEY]) == 1)
st2 = R.onboarding_intake_state(org_id=ORG)
row2 = next(r for r in st2["rail"]["verify_table"] if r["instance_key"] == IKEY)
check("D5 the Stage-4 row is now GREEN with the activation note '39 new activation, 23 bring-your-own-device, 25 upgrade'",
      row2["red"] is False and row2["status"] == "verified" and "39 new activation" in (row2["activation_note"] or "")
      and "25 upgrade" in row2["activation_note"] and "23 bring-your-own-device" in row2["activation_note"], row2.get("activation_note"))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§E every reader dereferences the saved rules")
use(db)
acfg = R._accessory_config(db, ORG)
lr = acfg["line_rules"]
check("E1 _accessory_config['line_rules'] resolves the saved JSON: source tenant, the fields and tokens, the basis key untouched",
      lr["source"] == "tenant" and lr["fields"] == prop["fields"] and lr["tokens"]["byod"] == prop["tokens"]["byod"]
      and acfg["activation_details_rules_raw"]["edge_name_tokens"] == ["edge plan"])
landed = [r for r in db.tables["raw_sales"] if r["org_id"] == ORG]
cells = R._sales_cell_agg(landed, acfg, exec_cfg=R._exec_metric_config(db, ORG))
prem = set().union(*(x["_prem"] for x in cells.values()))
upg = set().union(*(x["_upg"] for x in cells.values()))
byod = set().union(*(x["_byod"] for x in cells.values()))
check("E2 _sales_cell_agg over the LANDED rows (the Sales Report / Executive MTD / Daily Targets pass): 39 activations, 25 upgrades, 23 BYOD by distinct invoice, 67 phones",
      len(prem) == 39 and len(upg) == 25 and len(byod) == 23 and sum(x["total_phones"] for x in cells.values()) == 67,
      (len(prem), len(upg), len(byod), sum(x["total_phones"] for x in cells.values())))
try:
    ex = R._exec_mtd(db, ORG, "August 2026", today=date(2026, 9, 21))
    tot = ex["by_location"]["total"]
    cl = (ex.get("landing") or {}).get("classified") or {}
    check("E3 GET /exec-mtd (the real endpoint core): Total Activation 87 with the split, Total Phones 67; landing.classified gate CLOSED",
          int(tot.get("total_activation") or 0) == 87 and int(tot.get("upgrade") or 0) == 25 and int(tot.get("byod") or 0) == 23
          and int(tot.get("total_phones") or 0) == 67 and cl.get("gate_open") is False and cl.get("activation_type_transactions") == 87,
          ({k: tot.get(k) for k in ("total_activation", "activation", "port", "byod", "upgrade", "total_phones")}, cl.get("gate_open")))
except Exception as e:
    check("E3 GET /exec-mtd core ran over the fake client", False, repr(e)[:300])
check("E4 the pay path: the commission calc's cfg carries the org's rules (calc_rep_commissions reads cfg['line_class_rules'])",
      "line_class_rules" in open(os.path.join(os.path.dirname(__file__), "app/modules/commcalc/router.py")).read()
      and LC.classify_line(landed[0], lr) in ("premium", "upgrade", "byod", None))
sr = R.sales_report_detail if hasattr(R, "sales_report_detail") else None
check("E5 the mig-224 transaction rescue still SUPPLEMENTS (never overrides) the line predicate: a tid with a classified line is never rescued",
      R._blank_ct_bucket_map(landed, lr, [{"bucket": "premium", "all_of": [{"field": "category", "contains_any": ["features"]}]}]).keys()
      == {t for t in {r["trans_id"] for r in landed} if not any(LC.classify_line(r, lr) for r in landed if r["trans_id"] == t)
          and any("features" in str(r.get("category") or "").lower() for r in landed if r["trans_id"] == t)})

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§F the attestation path, and the Exec-MTD banner for an unmapped tenant")
db2 = fresh_db()
ACC_ROWS = [dict(r, category=CAT["acc"], product_desc="Phone Case", department="Accessories") for r in ROWS[:30]]
a2 = analyze(db2, sales_xlsx(ACC_ROWS), "acc.xlsx", source_kind="sales", pos_source="mypos", layout="sales",
             column_map=json.dumps(SALES_MAP), typed_total=f"{sum(r['ext_price'] for r in ACC_ROWS):.2f}")
check("F1 an accessory-only export: gate open at analyze, no proposal for any activation class (no hint hits → house tokens kept)",
      a2["line_class"]["gate_open"] is True and a2["line_class"]["suggest"]["proposal"]["tokens"]["activation"] == LC.HOUSE_TOKENS["activation"]
      and not a2["line_class"]["suggest"]["per_class"]["activation"])
c2 = commit(db2, sales_xlsx(ACC_ROWS), "acc.xlsx", source_kind="sales", pos_source="mypos", layout="sales",
            column_map=json.dumps(SALES_MAP), typed_total=f"{sum(r['ext_price'] for r in ACC_ROWS):.2f}")
check("F2 committed without an attestation → landed, not verified (needs_input)", c2["ok"] and not c2["verified"])
c3 = commit(db2, sales_xlsx(ACC_ROWS), "acc.xlsx", source_kind="sales", pos_source="mypos", layout="sales",
            column_map=json.dumps(SALES_MAP), typed_total=f"{sum(r['ext_price'] for r in ACC_ROWS):.2f}",
            attestation=json.dumps({"no_activations": "accessory-only kiosk, no lines of service sold here"}))
check("F3 committed WITH the attestation (a reason) → verified, the attestation recorded with the name",
      c3["ok"] and c3["verified"] and c3["verified_numbers"]["activation_gate"]["attested"]["by"] == "tester"
      and "kiosk" in c3["verified_numbers"]["activation_gate"]["attested"]["reason"], (c3.get("problems"), c3.get("activation_gate")))
try:
    ex2 = R._exec_mtd(db2, ORG, "August 2026", today=date(2026, 9, 21))
    cl2 = (ex2.get("landing") or {}).get("classified") or {}
    check("F4 the Executive MTD `landing.classified` block for an unmapped slice: gate OPEN, the note names the fields and the step, map_at → onboarding_intake / 2.5a",
          cl2.get("gate_open") is True and "2.5a" in (cl2.get("note") or "") and cl2["map_at"] == {"screen": "onboarding_intake", "step": "2.5a"}, cl2)
except Exception as e:
    check("F4 GET /exec-mtd core ran over the fake client (unmapped tenant)", False, repr(e)[:300])
p3 = R.onboarding_intake_put_line_class(R.OnboardingLineClassIn(instance_key=OI.stage2_instance_key("sales", "mypos", "sales"),
                                                                 no_activations="accessory-only kiosk", by="pat"), org_id=ORG)
check("F5 the 2.5a save can record the attestation on its own (no rule change): verified, attested by name",
      p3["ok"] and p3["verified"] is True and p3["activation_gate"]["attested"]["by"] in ("pat", "tester"), p3.get("activation_gate"))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("§G negative controls")
status, detail = http_error(R.onboarding_intake_put_line_class, R.OnboardingLineClassIn(instance_key=IKEY, fields="category"), org_id=ORG)
check("G1 fields that are not a list → 400", status == 400, (status, detail))
status, detail = http_error(R.onboarding_intake_put_line_class, R.OnboardingLineClassIn(instance_key=IKEY, tokens=["byod"]), org_id=ORG)
check("G2 tokens that are not an object → 400", status == 400, (status, detail))
status, detail = http_error(R.onboarding_intake_put_line_class, R.OnboardingLineClassIn(instance_key=IKEY, metric_rules={"activation": {"byod": ["x"]}}), org_id=ORG)
check("G3 a metric bucket that is not a LINE bucket ('activation' — its home is the predicate) → 400", status == 400, (status, detail))
status, detail = http_error(R.onboarding_intake_put_line_class, R.OnboardingLineClassIn(instance_key="sales:nope:sales", fields=["category"]), org_id=ORG)
check("G4 an unknown instance → 404 (after the config save, which is org-wide by design)", status == 404, (status, detail))
status, detail = http_error(R.put_accessory_config, R.PutAccessoryConfigIn(activation_details_rules=["x"]), ORG, "")
check("G5 PUT /accessory-config refuses a non-object activation_details_rules", status == 400, (status, detail))
db3 = FakeDB()
use(db3)
h = R._accessory_config(db3, HOUSE)
check("G6 an org with no config row → house rules (contract_type only, the retired tokens verbatim) — byte-identical",
      h["line_rules"]["source"] == "house" and h["line_rules"]["fields"] == ["contract_type"] and h["line_rules"]["tokens"] == LC.HOUSE_RULES["tokens"])
check("G7 the org-scope rule: every read the step makes is org-scoped (the fake filters by org_id; another org's rows never leak)",
      all(r["org_id"] == ORG for r in db.tables["accessory_config"]) and R._accessory_config(db, "other-org")["line_rules"]["source"] == "house")

print(f"\n══ onboarding intake — Stage D-2 (2.5a): {_pass} passed, {_fail} failed ══")
sys.exit(1 if _fail else 0)
