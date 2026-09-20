"""DB-FREE PROOF — the tenant-onboarding intake, STAGE C: the typed "other" reports land in their REAL
destinations, bill payments are EXTRACTED from the sales export, and the inventory listing auto-checks
itself against the ACTIVATION feed (owner 2026-09-20).

    python3 harness_onboarding_intake_c.py        (from backend/; no network, no DB)

OWNER'S WORDS (the acceptance bar): "other reports land in their respective categories if such a
report is present for that carrier, otherwise bill payments should be extracted from their sales
reports and then assigned a separate report for themselves. merchant payments are the merchant reports
which are used to reconcile the credit card payments received in the store … X report has their own
report where the data should be uploaded. Inventory also lands on the inventory module where it shows
inventory on hand. inventory report should also auto check itself with the activation report to see
which item has been sold but not rung out properly from the inventory."

THE SCENARIO, over an in-memory fake of the supabase client (+ its storage), driving the REAL router
functions — the endpoint's own paths (the X-report parser + `_xreport_land_rows`, the portal normalizer
+ `store_settlement`, `_ingest_mapped_df` into the processor feed, `epay_ingest.ingest`,
`b2b_sweep.write_inventory_devices`, the re-reads through the destinations' OWN readers):

    §B an X-report workbook — 2 stores (one the roster knows by address, one it does not), cash /
       card / other, a blank-label totals row per sheet → lands in pos_tender_summary, re-read ties
       per store-day, and the CLOSING recon's own `_addr_resolver` / `_xreport_tenders_by_store`
       resolves the landed store strings (one store map, not two)
    §C a merchant settlement export — 2 merchants × 2 days × card brands, fees, a TOTAL row → lands in
       merchant_settlement_day with the right role as an UPLOAD, re-read ties Σ gross / net / fees /
       count; the closing card recon reads it unchanged (`closing.router._settlement_rows_for_days`
       → `external_credit_recon.settlement_cells`); an overlap with the scheduled pull is REFUSED
    §D a carrier bill-pay report (the daily-tx feed layout) — lands through the mapped importer into
       the processor feed table, RE-READ through the mig-939 reader the coverage recon calls; and the
       transaction-detail layout through the feed's own idempotent ingest
    §E a sales export with bill-pay department lines → the extraction equals the ONE predicate's
       lines exactly (`exec_metric_defs.line_match`, and `_sales_cell_agg`'s bill_qty / bill_amt),
       with provenance per line; the Bill Payments report endpoint; beside the carrier feed
    §F an inventory listing + an activation feed with 3 activated-still-on-hand units — one by IMEI,
       one by mobile number THROUGH the sale line (a refunded sale), one unpairable (reported, never
       guessed) — and the negative control that an activation on a sales hit is EVIDENCE, never a
       second row
    §G Stage-4 rows carry the cross-check note; Stage-5 runbook lines name the real destinations
    §H NEGATIVE CONTROLS — the save guarantee per kind (a landing that drops rows → ok:false), a
       missing close date, a multi-day X-report, an unmapped merchant, a bad role
    §I wired, registered, RULE TWO (no carrier / POS / processor name in the intake code), the pages
"""
import asyncio
import inspect
import io
import json
import os
import re
import sys
import uuid

sys.path.insert(0, os.path.abspath(os.path.dirname(__file__)))

from app.modules.commcalc import onboarding_intake as OI
from app.modules.commcalc import column_mapping as CM
from app.modules.commcalc import billpay_extract as BPX
from app.modules.commcalc import inventory_sold_recon as ISR
from app.modules.commcalc import merchant_portals as MP
from app.modules.commcalc import exec_metric_defs as EMD
from app.modules.commcalc import epay_ingest as EPI
from app.modules.commcalc.device_cost_recon import device_key as DKEY
from app.modules.commcalc import router as R
from app.modules.storeops import router as SO
from app.modules.storeops import merchant_ids as MIDS
from app.modules.closing import router as CR
from app.modules.closing import external_credit_recon as ECR

ROOT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..")
MIG_DIR = os.path.join(ROOT, "database", "migrations")
FE_DIR = os.path.join(ROOT, "frontend", "src", "app", "(platform)")
FE_PAGE = os.path.join(FE_DIR, "onboarding", "intake", "page.tsx")
FE_STAGE2 = os.path.join(FE_DIR, "onboarding", "intake", "stage2.tsx")
FE_SHARED = os.path.join(FE_DIR, "onboarding", "intake", "intake-shared.tsx")
FE_BILLPAY = os.path.join(FE_DIR, "commcalc", "bill-payments", "page.tsx")
FE_INV = os.path.join(FE_DIR, "commcalc", "inventory-sold-recon", "page.tsx")
ORG = "11111111-2222-3333-4444-555555555555"
HOUSE = "00000000-0000-0000-0000-000000000001"

_pass = _fail = 0
_failures = []


def check(name, cond, extra=""):
    global _pass, _fail
    if cond:
        _pass += 1
        print(f"  PASS  {name}")
    else:
        _fail += 1
        _failures.append(name)
        print(f"  FAIL  {name}{(' — ' + str(extra)[:600]) if extra else ''}")
    return bool(cond)


def section(t):
    print(f"\n{t}\n" + "─" * len(t))


def money(x):
    return round(float(x or 0.0), 2)


# ══════════════════════════════════════════════════════════════════════════════════════════════════
# AN IN-MEMORY SUPABASE CLIENT (+ storage) — the harness_onboarding_intake_b pattern, extended with the
# Stage-C destination tables (declared columns; a select of an undeclared column raises like 42703)
# ══════════════════════════════════════════════════════════════════════════════════════════════════
class _Q:
    def __init__(self, db, table):
        self.db, self.table = db, table
        self.op, self.rows, self.filters, self.conflict = "select", None, [], None
        self._order, self._limit, self._range, self.cols = None, None, None, "*"
        self._count = None

    def select(self, cols="*", **kw):
        self.op, self.cols = "select", cols
        for c in [c.strip() for c in str(cols).split(",") if c.strip() and c.strip() != "*"]:
            if c not in self.db.columns(self.table):
                raise RuntimeError(f'42703 column "{c}" of {self.table} does not exist')
        if kw.get("count"):
            self._count = True
        return self

    def insert(self, rows):
        self.op, self.rows = "insert", (rows if isinstance(rows, list) else [rows])
        return self

    def upsert(self, rows, on_conflict=None):
        self.op, self.rows, self.conflict = "upsert", (rows if isinstance(rows, list) else [rows]), on_conflict
        return self

    def update(self, row, count=None):
        self.op, self.rows, self._count = "update", [row], count
        return self

    def delete(self):
        self.op = "delete"
        return self

    def eq(self, k, v):    self.filters.append(("eq", k, v)); return self
    def neq(self, k, v):   self.filters.append(("neq", k, v)); return self
    def in_(self, k, v):   self.filters.append(("in", k, list(v))); return self
    def is_(self, k, v):   self.filters.append(("is", k, v)); return self
    def gte(self, k, v):   self.filters.append(("gte", k, v)); return self
    def lte(self, k, v):   self.filters.append(("lte", k, v)); return self
    def lt(self, k, v):    self.filters.append(("lt", k, v)); return self
    def ilike(self, k, v): self.filters.append(("ilike", k, v)); return self
    def order(self, k, desc=False): self._order = (k, desc); return self
    def limit(self, n):    self._limit = n; return self
    def range(self, lo, hi): self._range = (lo, hi); return self

    def _match(self, r):
        for op, k, v in self.filters:
            x = r.get(k)
            if op == "eq" and x != v: return False
            if op == "neq" and x == v: return False
            if op == "in" and x not in v: return False
            if op == "is" and not ((v == "null" and x is None) or (v is not None and v != "null" and x == v)): return False
            if op == "gte" and not (x is not None and str(x) >= str(v)): return False
            if op == "lte" and not (x is not None and str(x) <= str(v)): return False
            if op == "lt" and not (x is not None and str(x) < str(v)): return False
            if op == "ilike" and not (x is not None and str(v).strip("%").lower() in str(x).lower()): return False
        return True

    def execute(self):
        t = self.db.tables.setdefault(self.table, [])
        if t is None:
            raise RuntimeError(f'42P01 relation "{self.table}" does not exist')
        if self.op == "select":
            out = [dict(r) for r in t if self._match(r)]
            if self._order:
                k, desc = self._order
                out.sort(key=lambda r: (r.get(k) is None, str(r.get(k))), reverse=desc)
            n_all = len(out)
            if self._range:
                lo, hi = self._range
                out = out[lo:hi + 1]
            if self._limit is not None:
                out = out[:self._limit]
            return _Res(out, count=n_all if self._count else None)
        if self.op == "insert":
            self.db.check_rows(self.table, self.rows)
            written = []
            for r in self.rows:
                row = dict(r)
                row.setdefault("id", str(uuid.uuid4()))
                if self.db.drop_inserts.get(self.table):
                    continue
                if self.db.drop_every_other.get(self.table) and len(written) % 2:
                    written.append(dict(row))
                    continue
                t.append(row)
                written.append(dict(row))
            return _Res(written)
        if self.op == "upsert":
            self.db.check_rows(self.table, self.rows, upsert=True)
            keys = [k.strip() for k in (self.conflict or "").split(",") if k.strip()]
            written = []
            for r in self.rows:
                row = dict(r)
                hit = next((x for x in t if keys and all(x.get(k) == row.get(k) for k in keys)
                            and all(row.get(k) is not None for k in keys)), None)
                if hit:
                    hit.update(row)
                    written.append(dict(hit))
                else:
                    row.setdefault("id", str(uuid.uuid4()))
                    if not self.db.drop_inserts.get(self.table):
                        t.append(row)
                    written.append(dict(row))
            return _Res(written)
        if self.op == "update":
            self.db.check_rows(self.table, self.rows)
            out = []
            for r in t:
                if self._match(r):
                    r.update(self.rows[0])
                    out.append(dict(r))
            return _Res(out, count=len(out) if self._count else None)
        if self.op == "delete":
            keep = [r for r in t if not self._match(r)]
            gone = len(t) - len(keep)
            t[:] = keep
            return _Res([{}] * gone)
        raise RuntimeError(self.op)


class _Res:
    def __init__(self, data, count=None): self.data, self.count = data, count


class FakeStorage:
    def __init__(self, broken=False):
        self.buckets, self.files, self.broken = set(), {}, broken

    def get_bucket(self, name):
        if name not in self.buckets:
            raise RuntimeError("bucket not found")
        return {"name": name}

    def create_bucket(self, name):
        if self.broken:
            raise RuntimeError("storage unavailable")
        self.buckets.add(name)

    def from_(self, bucket):
        st = self

        class _B:
            def upload(self_, path, data, opts=None):
                if st.broken:
                    raise RuntimeError("storage unavailable")
                st.files[(bucket, path)] = bytes(data)
                return {"Key": path}

            def download(self_, path):
                if (bucket, path) not in st.files:
                    raise RuntimeError("object not found")
                return st.files[(bucket, path)]
        return _B()


class FakeDB:
    def __init__(self, storage_broken=False):
        self.tables = {}
        self.drop_inserts = {}
        self.drop_every_other = {}
        self.storage = FakeStorage(broken=storage_broken)
        self.declared = {
            "carrier": ["id", "org_id", "name", "code", "is_default"],
            "column_mapping": ["id", "org_id", "report_key", "carrier_id", "target_field", "source_header",
                               "transform", "is_active", "priority", "updated_at", "sign_convention"],
            # `source` = mig 727 (writer provenance) — since 2026-09-20 also the report KIND that wrote the row
            # (landing_identity.KIND_STAMP); `raw_sales_product` = mig 1011, the by-product aggregate's OWN table
            "raw_sales": ["id", "org_id", "period", "period_month", "period_year", "store", "salesperson", "user_login",
                          "department", "category", "product_desc", "product_id", "gp", "ext_price", "trans_id",
                          "trans_date", "contract_type", "mdn", "serial_1", "register", "tender_type", "voided",
                          "trans_type", "sku", "customer", "email", "customer_no", "quantity", "total_cost",
                          "pricing_discounts", "contract_no", "source", "created_at"],
            "raw_sales_product": ["id", "org_id", "period", "period_month", "period_year", "store", "salesperson", "user_login",
                                  "department", "category", "product_desc", "product_id", "gp", "ext_price", "trans_id",
                                  "trans_date", "contract_type", "mdn", "serial_1", "register", "tender_type", "voided",
                                  "trans_type", "sku", "customer", "email", "customer_no", "quantity", "total_cost",
                                  "pricing_discounts", "contract_no", "source", "created_at"],
            "daily_sales_feed": ["id", "org_id", "period", "store", "salesperson", "department", "category", "product_desc",
                                 "gp", "ext_price", "trans_id", "trans_date", "contract_type", "mdn", "serial_1", "tender_type",
                                 "voided", "trans_type", "quantity"],
            "inventory_aging_device": ["id", "org_id", "imei", "serial", "sku", "item", "store", "unit_cost",
                                       "received_date", "days_in_stock", "as_of_date", "source", "raw_row",
                                       "on_hand", "off_hand_as_of", "status", "quantity", "total_cost", "category",
                                       "updated_at", "created_at"],
            "store_mapping": ["id", "org_id", "store_code", "store_address", "market", "is_active"],
            "stores": ["id", "org_id", "store_code", "address", "market", "is_active", "entity_id", "timezone"],
            "store_aliases": ["id", "org_id", "alias", "store_code", "note", "source", "confidence"],
            "rep_aliases": ["id", "org_id", "alias", "canonical"],
            "employees": ["id", "org_id", "employee_id", "name", "home_store", "epay_salesperson", "is_active"],
            "onboarding_run": ["id", "org_id", "run_kind", "period", "status", "current_step", "started_by",
                               "started_at", "signed_off_by", "signed_off_on_behalf", "signed_off_at", "updated_at"],
            "onboarding_stage_state": ["id", "run_id", "org_id", "stage", "step", "instance_key", "status", "payload",
                                       "verified_numbers", "verified_by", "verified_at", "blocking_reason", "updated_at"],
            # ── Stage C destinations (each an EXISTING table; columns per their migrations) ──
            "pos_tender_summary": ["id", "org_id", "close_date", "store", "tender_type", "tender_class", "amount",
                                   "source", "updated_at"],                                             # mig 062
            "merchant_settlement_day": ["id", "org_id", "source_id", "portal_key", "report_key", "settlement_role",
                                        "business_date", "merchant_id", "terminal_id", "store_label", "store_code",
                                        "card_brand", "gross_amount", "refund_amount", "net_amount", "fee_amount",
                                        "txn_count", "currency", "batch_ref", "source_line", "raw", "pulled_at",
                                        "created_at"],                                                   # mig 955
            "raw_ma_daily_tx": ["id", "org_id", "period", "period_month", "period_year", "source_id", "tx_date",
                                "due_date", "account_id", "account_name", "direct_ma_id", "direct_ma_name",
                                "top_ma_id", "top_ma_name", "order_number", "user_name", "order_type",
                                "product_name", "retail_cost", "merchant_discount", "merchant_invoice"],   # mig 083
            "raw_epay_daily_tx": ["id", "org_id", "transaction_id", "transaction_source_id", "invoice_id",
                                  "settlement_date", "terminal_id", "user_name", "product", "product_title",
                                  "tx_type", "host_timestamp", "control_number", "retail", "discount", "cost",
                                  "commission", "is_fee", "source_batch", "store_code"],                # mig 903
            "store_merchant_id": ["id", "org_id", "store_code", "processor", "merchant_id", "not_required", "note",
                                  "created_at", "updated_at"],                                          # mig 902
            "data_source": ["id", "org_id", "processor", "label", "settlement_role", "portal_calibration",
                            "account_id", "enabled"],                                                   # mig 083/955
            "report_pull_map": ["id", "org_id", "report_key", "display_name", "target_table", "column_map",
                                "param_spec", "export_pref", "enabled", "sort_order", "processor"],       # mig 207
            "raw_custom_import": ["id", "org_id", "report_key", "period", "period_month", "period_year",
                                  "source_filename", "row_index", "data", "carrier_id"],                # mig 099
        }

    def columns(self, table):
        if table in self.declared:
            return self.declared[table]
        return ["*"] + [k for r in (self.tables.get(table) or []) for k in r]

    def check_rows(self, table, rows, upsert=False):
        if table in self.declared:
            for r in rows:
                bad = [k for k in r if k not in self.declared[table]]
                if bad:
                    raise RuntimeError(f'42703 column "{bad[0]}" of {table} does not exist')
        if table in ("raw_sales", "raw_sales_product", "raw_ma_daily_tx"):
            for r in rows:
                if not r.get("period"):
                    raise RuntimeError(f'23502 null value in column "period" of {table} violates not-null constraint')

    def schema(self, _name): return self
    def table(self, name): return _Q(self, name)

    def seed(self, table, rows):
        for r in rows:
            row = dict(r); row.setdefault("id", str(uuid.uuid4()))
            self.tables.setdefault(table, []).append(row)


class FakeUpload:
    def __init__(self, data, filename="export.xlsx"):
        self.data, self.filename = data, filename

    async def read(self): return self.data


STORE_ADDR = "100 Main St"          # the roster address of B-100
STORE_NEW = "Riverside Kiosk"       # unknown to the platform
PORTAL = MP.PORTAL_KEYS[0]          # a portal of the house registry — a DATA value, never named in code
MID_A, MID_B = "MID10001", "MID10002"


def fresh_db():
    db = FakeDB()
    db.seed("stores", [{"org_id": ORG, "store_code": "B-100", "address": STORE_ADDR, "market": "NY", "is_active": True},
                       {"org_id": ORG, "store_code": "B-200", "address": "200 Oak Ave", "market": "NY", "is_active": True}])
    db.seed("store_mapping", [{"org_id": ORG, "store_code": "B-100", "store_address": STORE_ADDR, "market": "NY"},
                              {"org_id": ORG, "store_code": "B-200", "store_address": "200 Oak Ave", "market": "NY"}])
    db.seed("employees", [{"org_id": ORG, "employee_id": "E1", "name": "Alice Rep", "epay_salesperson": "arep", "is_active": True},
                          {"org_id": ORG, "employee_id": "E2", "name": "Bob Smith", "epay_salesperson": "bob.s", "is_active": True}])
    db.seed("carrier", [{"id": "aaaaaaaa-0000-0000-0000-00000000c001", "org_id": ORG, "name": "Northwind Cellular", "code": "northwind"}])
    # the mig-902 merchant map: MID_A / ACCT-1 known, MID_B / ACCT-2 not yet
    db.seed("store_merchant_id", [{"org_id": ORG, "store_code": "B-100", "processor": PORTAL, "merchant_id": MID_A, "not_required": False},
                                  {"org_id": ORG, "store_code": "B-100", "processor": "vidapay", "merchant_id": "ACCT-1", "not_required": False}])
    # the mig-955 HOUSE registry row the closing card recon resolves the settlement feed through
    db.seed("report_pull_map", [{"org_id": HOUSE, "report_key": "merchant_settlement", "display_name": "Merchant portal settlement",
                                 "target_table": "commcalc.merchant_settlement_day",
                                 "column_map": {"day": "business_date", "amount": "net_amount", "role": "settlement_role",
                                                "store_code": "store_code", "merchant_id": "merchant_id"},
                                 "enabled": True, "processor": "merchant_portal"}])
    return db


def use(db):
    R.sb = lambda: db
    SO.sb = lambda: db
    SO.get_supabase = lambda: db
    MIDS.get_supabase = lambda: db
    CR.get_supabase = lambda: db
    EPI.get_supabase = lambda: db
    R._TABLE_COL_PRESENT.clear()
    return db


def run(coro):
    return asyncio.run(coro)


def analyze(db, data, filename, **form):
    use(db)
    kw = dict(org_id=ORG)
    kw.update(form)
    return run(R.onboarding_intake_analyze(file=FakeUpload(data, filename), **kw))


def commit(db, data, filename, **form):
    use(db)
    kw = dict(verified_by="tester", org_id=ORG)
    kw.update(form)
    return run(R.onboarding_intake_commit(file=FakeUpload(data, filename), **kw))


def http_error(fn, *a, **kw):
    try:
        fn(*a, **kw)
        return None, None
    except R.HTTPException as e:
        return e.status_code, e.detail


# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("A. THE VOCABULARY (pure) — known kinds, real destinations, the per-kind numbers")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
check("the 2.0 picker offers the three KNOWN kinds + free text; each known kind names an EXISTING destination table",
      OI.OTHER_KINDS == ("x_report", "merchant_payments", "bill_payments", "other")
      and OI.SOURCE_KIND_TARGET["x_report"] == "pos_tender_summary" and OI.SOURCE_KIND_TARGET["merchant_payments"] == "merchant_settlement_day"
      and OI.SOURCE_KIND_TARGET["bill_payments"] in OI.BILLPAY_FEED_TABLES and "other" not in OI.SOURCE_KIND_TARGET
      and all(k in OI.OTHER_KIND_HINTS for k in OI.OTHER_KINDS))
check("the matrix kinds map through their destination's OWN parser — no column layout; a bill-pay report offers BOTH processor feed layouts, the daily-tx feed first",
      OI.layouts_for_kind("x_report", CM.TABLE_MAP) == [] and OI.layouts_for_kind("merchant_payments", CM.TABLE_MAP) == []
      and [l["report_key"] for l in OI.layouts_for_kind("bill_payments", CM.TABLE_MAP)] == ["ma_daily_tx", "epay_daily_tx"]
      and {l["target_table"] for l in OI.layouts_for_kind("bill_payments", CM.TABLE_MAP)} == set(OI.BILLPAY_FEED_TABLES))
check("stage-2 instance keys for the new kinds; every one is a stage-2 kind",
      OI.stage2_instance_key("x_report", "rq", "x_report") == "x_report:rq:x_report"
      and OI.stage2_instance_key("merchant_payments", PORTAL, "settlement").startswith("merchant_payments:")
      and all(OI.stage_of_kind(k) == "2" for k in OI.OTHER_KINDS))
check("a bill-pay report's fields follow its LAYOUT (the processor's account / terminal id is the store; the mig-902 map resolves it)",
      OI.kind_fields("bill_payments", "ma_daily_tx")["store"] == "account_id" and OI.kind_fields("bill_payments", "epay_daily_tx")["store"] == "terminal_id"
      and OI.kind_required("bill_payments", "ma_daily_tx") == ("account_id", "tx_date", "retail_cost", "product_name"))
check("the X-report close date comes from the FILENAME (a single day); a range is refused; no date is None — never today",
      OI.xreport_file_date("X-Report_09012026-09012026.xlsx") == "2026-09-01"
      and OI.xreport_file_date("X-Report_09012026-09022026.xlsx") == {"range": True, "from": "2026-09-01", "to": "2026-09-02"}
      and OI.xreport_file_date("xreport.xlsx") is None)
_xrows = [(STORE_ADDR, "2026-09-01", "Cash", 522.08), (STORE_ADDR, "2026-09-01", "Credit Card", 112.22), (STORE_ADDR, "2026-09-01", "Check", 10.00),
          (STORE_NEW, "2026-09-01", "Cash", 100.00), (STORE_NEW, "2026-09-01", "Debit", 50.00)]
_xv = OI.xreport_verify(_xrows, R._xr_tender_class, {STORE_ADDR: 644.30})
check("xreport_verify: cash / card / other per store-day by the import's ONE class rule, beside each sheet's own total",
      _xv["sum_cash"] == 622.08 and _xv["sum_card"] == 162.22 and _xv["sum_other"] == 10.0 and _xv["sum_total"] == 794.30
      and _xv["per_store_day"][0]["file_total"] == 644.30 and _xv["per_store_day"][0]["difference"] == 0.0
      and _xv["per_store_day"][1]["file_total"] is None and _xv["file_total"] == 644.30, _xv)
check("xreport_sheet_total: the blank-label totals row's Net cell, after the header",
      OI.xreport_sheet_total([["Tender Types", "Sales", "Net"], ["Cash", "1", "1.00"], ["", "$2.50", "$2.50"]], 0, 2) == 2.5
      and OI.xreport_sheet_total([["Tender Types", "Sales", "Net"], ["Cash", "1", "1.00"], ["Total", "", "1.00"]], 0, 2) == 1.0
      and OI.xreport_sheet_total([["Tender Types", "Sales", "Net"], ["Cash", "1", "1.00"]], 0, 2) is None)
_mv = OI.merchant_verify([{"merchant_id": "M1", "business_date": "2026-09-01", "gross_amount": 100.0, "net_amount": 95.0, "fee_amount": 2.5, "txn_count": 3, "card_brand": "visa", "store_code": "B-100"},
                          {"merchant_id": "M1", "business_date": "2026-09-01", "gross_amount": 50.0, "net_amount": 50.0, "fee_amount": 1.0, "txn_count": 1, "card_brand": "amex", "store_code": "B-100"},
                          {"merchant_id": "M2", "business_date": "2026-09-02", "gross_amount": 10.0, "net_amount": 10.0, "fee_amount": 0.5, "txn_count": 1, "card_brand": "visa", "store_code": None}],
                         {"gross": 160.0, "net": 155.0, "fees": 4.0, "count": 5})
check("merchant_verify: Σ gross / net / fees / count per merchant per business day beside the file's TOTAL row; an unmapped merchant is listed, never $0",
      _mv["sum_gross"] == 160.0 and _mv["sum_net"] == 155.0 and _mv["sum_fees"] == 4.0 and _mv["txn_count"] == 5
      and len(_mv["per_merchant_day"]) == 2 and _mv["difference"] == {"gross": 0.0, "net": 0.0, "fees": 0.0, "count": 0}
      and _mv["unmapped_merchants"] == ["M2"] and _mv["by_brand"]["visa"] == 105.0, _mv)
_bv = OI.billpay_feed_verify([{"account_id": "A1", "tx_date": "2026-09-01", "retail_cost": "50", "order_type": "Sales Order", "product_name": "RTR Payment"},
                              {"account_id": "A1", "tx_date": "2026-09-01", "retail_cost": "599.99", "order_type": "Sales Order", "product_name": "Handset"},
                              {"account_id": "A2", "tx_date": "2026-09-02", "retail_cost": "25", "order_type": "Sales Order", "product_name": "RTR Payment"}],
                             OI.BILLPAY_LAYOUT_FIELDS["ma_daily_tx"], lambda m: "rtr" in str(m.get("product_name")).lower(),
                             resolve_store=lambda a: {"A1": "B-100"}.get(a, ""))
check("billpay_feed_verify: every row counted, the org's bill-pay rows Σ per store-day, the unmapped account listed",
      _bv["rows"] == 3 and _bv["billpay_rows"] == 2 and _bv["sum_all_rows"] == 674.99 and _bv["sum_amount"] == 75.0
      and _bv["unmapped_accounts"] == ["A2"] and _bv["per_store_day"][0]["store"] == "B-100", _bv)
check("stage4_note says what the cross-check found, in the owner's words",
      "activated but still on hand" in OI.stage4_note({"sold_check": {"basis": "x", "activated_not_rung_out": 3, "sold_not_cleared": 1, "activations_present": True}})
      and "bill-pay line(s) extracted" in OI.stage4_note({"billpay_extract": {"basis": "x", "lines": 4, "sum": 100.0, "feed_present": False}})
      and OI.stage4_note({}) is None)
check("lands_in: the kind's destination, or the layout's table a bill-pay commit recorded; free text has none",
      OI.lands_in("x_report") == "pos_tender_summary" and OI.lands_in("bill_payments", {"target_table": "raw_epay_daily_tx"}) == "raw_epay_daily_tx"
      and OI.lands_in("other").startswith("(no destination"))

# the extraction predicate is line_match's answer, plus WHY — pinned over a grid of cases
_rule = {"department": ["rtr"], "category": ["rtr product", "other carr. payments"], "exclude_category": ["other charge"],
         "product_desc_contains": ["refill"]}
_grid = [("rtr", "other charge", "boost rtr payment"), ("rtr", "rtr product", "x"), ("phones", "rtr product", "x"),
         ("phones", "cases", "xfinity refill"), ("phones", "cases", "case"), ("", "", ""), ("rtr", "", "")]
check("match_reason ≡ exec_metric_defs.line_match on every case (exclusions first, then category, department, product) — plus the token that matched",
      all((BPX.match_reason(_rule, d, c, p) is not None) == EMD.line_match(_rule, d, c, p) for d, c, p in _grid)
      and BPX.match_reason(_rule, "rtr", "rtr product", "x") == {"by": "category", "token": "rtr product"}
      and BPX.match_reason(_rule, "rtr", "", "") == {"by": "department", "token": "rtr"}
      and BPX.match_reason(_rule, "phones", "cases", "xfinity refill") == {"by": "product_desc_contains", "token": "refill"}
      and BPX.match_reason(_rule, "rtr", "other charge", "boost rtr payment") is None)
check("is_countable mirrors _sales_cell_agg's three skip rules (voided / Return / no rep or admin)",
      not BPX.is_countable({"voided": "Yes", "salesperson": "a"}) and not BPX.is_countable({"trans_type": "Return", "salesperson": "a"})
      and not BPX.is_countable({"salesperson": "admin"}) and not BPX.is_countable({"salesperson": ""}) and BPX.is_countable({"salesperson": "Alice"}))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("B. THE X-REPORT — through the existing parser, into pos_tender_summary, re-read; ONE store map with the closing recon")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
XR_HEADER = ["Tender Types", "Sales", "Returns/Trade In", "Sub Net", "Drop", "Pickup", "Payments", "Refunds", "Net"]


def xr_money_row(label, amt):
    return [label, f"{amt:.2f}", "0.00", f"{amt:.2f}", "0.00", "0.00", "0.00", "0.00", f"{amt:.2f}"]


def xr_sheet(store, tenders, total=True, extra_label=None):
    rows = [[f"Report Date:9/1/2026 - 9/1/2026", "", f"Locations:{store}", "", "", "", "", "", ""],
            ["Report Grouping:No grouping", "", "", "", "", "", "", "", ""],
            ["Tendered Amounts", "", "", "", "", "", "", "", ""], XR_HEADER]
    rows += [xr_money_row(l, a) for l, a in tenders]
    if extra_label:
        rows.append(xr_money_row(extra_label, 7.77))
    if total:
        t = sum(a for _l, a in tenders) + (7.77 if extra_label else 0)
        rows.append(["", f"${t:,.2f}", "$0.00", f"${t:,.2f}", "$0.00", "$0.00", "$0.00", "$0.00", f"${t:,.2f}"])
    return rows


XR_A = [("Cash", 522.08), ("Check", 0.00), ("Credit Card", 112.22), ("Gift Card", 0.00), ("Store Account", 15.50)]
XR_B = [("Cash", 100.00), ("Debit Card", 50.00), ("Check", 5.00)]
XR_TOTAL = money(sum(a for _l, a in XR_A) + sum(a for _l, a in XR_B))


def xr_workbook(sheets):
    import pandas as pd
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as w:
        for name, rows in sheets:
            pd.DataFrame(rows).to_excel(w, sheet_name=name, header=False, index=False)
    return buf.getvalue()


XR_FILE = xr_workbook([(STORE_ADDR, xr_sheet(STORE_ADDR, XR_A)), (STORE_NEW, xr_sheet(STORE_NEW, XR_B))])
XR_NAME = "X-Report_09012026-09012026.xlsx"

db = fresh_db()
x1 = analyze(db, XR_FILE, XR_NAME, source_kind="x_report", pos_source="rq")
xv = x1["verify"]["numbers"]
check("2.2/2.5: both sheets parsed by the existing X-report parser; per store-day cash / card / other; the sheets' own totals beside ours, to the cent",
      x1["source_kind"] == "x_report" and x1["matrix"] is True and x1["target_table"] == "pos_tender_summary"
      and xv["close_date"] == "2026-09-01" and xv["date_source"] == "filename" and xv["stores"] == 2
      and xv["sum_cash"] == 622.08 and xv["sum_card"] == 162.22 and xv["sum_other"] == 20.50 and xv["sum_total"] == XR_TOTAL
      and xv["file_total"] == XR_TOTAL and x1["verify"]["tie"]["match"] is True and xv["parser"]["headers_found"] == 2, xv)
st = {r["value"]: r for r in x1["stores"]}
check("2.4: the sheet names ARE the store strings — the roster address resolves (address), the kiosk is UNRESOLVED and named in the refusals",
      st[STORE_ADDR]["lands_as"] == "B-100" and st[STORE_ADDR]["how"] == "address" and st[STORE_NEW]["status"] == "unresolved"
      and any("unresolved" in r for r in x1["verify"]["refusals"]) and x1["unresolved_stores"] == [STORE_NEW])
s_, d_ = http_error(commit, db, XR_FILE, XR_NAME, source_kind="x_report", pos_source="rq")
check("COMMIT refused while the kiosk is unresolved — nothing landed", s_ == 400 and "unresolved" in str(d_) and not db.tables.get("pos_tender_summary"))
DEC_X = {"store": {STORE_NEW: {"action": "assign", "store_code": "B-200"}}}
c1 = commit(db, XR_FILE, XR_NAME, source_kind="x_report", pos_source="rq", identity=json.dumps(DEC_X))
pts = db.tables["pos_tender_summary"]
check("COMMIT ok: 8 tender rows landed through _xreport_land_rows (mig-062 key, ONE class rule, source stamped), the alias written for the kiosk",
      c1["ok"] is True and c1["saved"] == 8 and len(pts) == 8 and all(r["source"] == "x_report" and r["close_date"] == "2026-09-01" for r in pts)
      and {r["tender_class"] for r in pts} == {"cash", "card", "other"} and c1["identity_written"]["aliases"] == [(STORE_NEW, "B-200")], c1.get("problems"))
cv = c1["verified_numbers"]["numbers"]
check("…RE-READ from pos_tender_summary: rows = built = inserted; Σ cash / card / other / total equal what was shown; ties the sheets' totals",
      c1["verified_numbers"]["rows_landed"] == 8 == c1["verified_numbers"]["rows_built"] and cv["sum_cash"] == 622.08 and cv["sum_card"] == 162.22
      and cv["sum_total"] == XR_TOTAL and c1["verified_numbers"]["tie"]["match"] is True
      and c1["verified_numbers"]["basis"] == "re-read from commcalc.pos_tender_summary after landing", c1["verified_numbers"])
use(db)
resolve_closing = CR._addr_resolver(db, ORG)
xt = CR._xreport_tenders_by_store(db, ORG, "2026-09-01")
check("ONE STORE MAP: the CLOSING recon's own _addr_resolver resolves both landed store strings (the kiosk through the alias the intake wrote) — no second map",
      resolve_closing(STORE_ADDR) == "B-100" and resolve_closing(STORE_NEW) == "B-200")
check("…and _xreport_tenders_by_store (every closing cash / card recon) reads the landed day: B-100 cash 522.08 / card 112.22, B-200 cash 100 / card 50",
      xt["B-100"]["cash"] == 522.08 and xt["B-100"]["card"] == 112.22 and xt["B-100"]["other"] == 15.50
      and xt["B-200"]["cash"] == 100.0 and xt["B-200"]["card"] == 50.0, xt)
c1b = commit(db, XR_FILE, XR_NAME, source_kind="x_report", pos_source="rq", identity=json.dumps(DEC_X))
check("re-commit is an UPSERT on the natural key — still 8 rows, never doubled", c1b["ok"] is True and len(db.tables["pos_tender_summary"]) == 8)
row_x = next(r for r in db.tables["onboarding_stage_state"] if r["instance_key"] == "x_report:rq:x_report")
check("the stage row is VERIFIED with the close date as its period and the destination recorded",
      row_x["status"] == "verified" and row_x["payload"]["period"] == "2026-09-01" and row_x["payload"]["target_table"] == "pos_tender_summary")

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("C. THE MERCHANT SETTLEMENT — the portal normalizer, mig 955 as an UPLOAD, the right role; the closing card recon reads it")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
MER_CSV = f"""Payments Hub - Transactions by Card Type
Generated 09/03/2026

Business Date,Merchant ID,DBA,Card Type,Sales Amount,Refunds,Net Amount,Fees,Transaction Count
09/01/2026,{MID_A},STORE ONE,Visa,"$1,240.50",(40.00),"$1,200.50",$32.15,18
09/01/2026,{MID_A},STORE ONE,MasterCard,$860.25,0.00,$860.25,$22.40,11
09/02/2026,{MID_A},STORE ONE,Visa,$980.00,(15.50),$964.50,$25.00,14
09/01/2026,{MID_B},STORE TWO,Visa,$540.00,0.00,$540.00,$14.00,7
09/02/2026,{MID_B},STORE TWO,Amex,$310.00,0.00,$310.00,$12.10,3
TOTAL,,,,"$3,930.75",(55.50),"$3,875.25",$105.65,53
""".encode("utf-8")
MER_NET, MER_GROSS, MER_FEES, MER_COUNT = 3875.25, 3930.75, 105.65, 53

db = fresh_db()
m1 = analyze(db, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL)
mv = m1["verify"]["numbers"]
check("2.2/2.5: the export read by the portal normalizer (banner skipped, TOTAL row skipped-and-used), Σ gross / net / fees / count per merchant-day beside the file's totals",
      m1["matrix"] is True and m1["target_table"] == "merchant_settlement_day" and mv["rows"] == 5 and len(mv["per_merchant_day"]) == 4
      and mv["sum_net"] == MER_NET and mv["sum_gross"] == MER_GROSS and mv["sum_fees"] == MER_FEES and mv["txn_count"] == MER_COUNT
      and mv["file_totals"]["net"] == MER_NET and mv["difference"] == {"gross": 0.0, "net": 0.0, "fees": 0.0, "count": 0}
      and m1["verify"]["tie"]["match"] is True, mv)
check("the role is the portal's house default (external_cc — the standalone terminal), stated with its source; the grain is settlement, never funding",
      m1["settlement_role"] == MP.settlement_role(PORTAL) and mv["role_source"] == "the portal's house default"
      and "never summed with funding" in mv["grain"] and mv["report_key"] == "intake_upload")
st = {r["value"]: r for r in m1["stores"]}
check("2.4: merchant ids through the mig-902 map — MID_A → B-100 (merchant_id), MID_B unresolved (its DBA 'STORE TWO' is no address the roster knows)",
      st[MID_A]["lands_as"] == "B-100" and st[MID_A]["how"] == "merchant_id" and st[MID_B]["status"] == "unresolved"
      and m1["identity_kind"] == "merchant_id" and st[MID_B].get("label") == "STORE TWO")
s_, d_ = http_error(commit, db, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL)
check("COMMIT refused while MID_B is unmapped — an unmapped merchant is never counted as a store's $0", s_ == 400 and MID_B in str(d_))
DEC_M = {"store": {MID_B: {"action": "assign", "store_code": "B-200"}}}
c2 = commit(db, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, identity=json.dumps(DEC_M))
msd = db.tables["merchant_settlement_day"]
sid = c2["verified_numbers"]["numbers"]["source_id"]
check("COMMIT ok: 5 rows landed through merchant_portal_sweep.store_settlement (mig-955 key), role external_cc, report_key intake_upload, a deterministic upload source_id",
      c2["ok"] is True and c2["saved"] == 5 and len(msd) == 5 and all(r["settlement_role"] == "external_cc" and r["report_key"] == "intake_upload"
                                                                     and r["source_id"] == sid and r["portal_key"] == PORTAL for r in msd)
      and sid == str(uuid.uuid5(uuid.NAMESPACE_URL, f"metricspro:onboarding-intake:{ORG}:{PORTAL}:external_cc")), c2.get("problems"))
check("…the MID_B decision wrote a storeops.store_merchant_id row (the SAME map the sweep resolves through) and read back; store_code stamped on every row",
      c2["identity_written"]["merchant_ids"] == [(MID_B, "B-200")]
      and any(r["merchant_id"] == MID_B and r["store_code"] == "B-200" and r["processor"] == PORTAL for r in db.tables["store_merchant_id"])
      and {r["store_code"] for r in msd} == {"B-100", "B-200"})
cv = c2["verified_numbers"]["numbers"]
check("…RE-READ from merchant_settlement_day: Σ gross / net / fees / count per merchant-day equal what was shown, to the cent",
      c2["verified_numbers"]["rows_landed"] == 5 and cv["sum_net"] == MER_NET and cv["sum_gross"] == MER_GROSS and cv["sum_fees"] == MER_FEES
      and cv["txn_count"] == MER_COUNT and c2["verified_numbers"]["tie"]["match"] is True)
use(db)
srows, days, spec = CR._settlement_rows_for_days(db, ORG, "2026-09-01", "2026-09-02")
cells, unmapped = ECR.settlement_cells(srows, MIDS.resolve_map(ORG, PORTAL))
check("THE CLOSING CARD RECON READS IT UNCHANGED: _settlement_rows_for_days resolves the feed through the mig-955 registry row and settlement_cells tallies per (store, day, role)",
      spec is not None and spec["table"] == "merchant_settlement_day" and len(srows) == 5 and unmapped == []
      and cells[("B-100", "2026-09-01")]["external_cc"] == 2060.75 and cells[("B-100", "2026-09-02")]["external_cc"] == 964.50
      and cells[("B-200", "2026-09-01")]["external_cc"] == 540.0 and cells[("B-200", "2026-09-02")]["external_cc"] == 310.0
      and days == {"external_cc": {"2026-09-01", "2026-09-02"}}, (cells, unmapped))
c2b = commit(db, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, identity=json.dumps(DEC_M))
check("re-commit restates the same (source, merchant, day, brand) rows in place — still 5", c2b["ok"] is True and len(db.tables["merchant_settlement_day"]) == 5)
# the OTHER role, by the person's pick
db2 = fresh_db()
m2 = analyze(db2, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, role="pos_merchant",
             identity=json.dumps(DEC_M))
c3 = commit(db2, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, role="pos_merchant", identity=json.dumps(DEC_M))
check("the person may state the OTHER role (the POS's own card tender) — rows land with it, under their own upload source id",
      m2["settlement_role"] == "pos_merchant" and m2["verify"]["numbers"]["role_source"] == "your pick" and c3["ok"] is True
      and all(r["settlement_role"] == "pos_merchant" for r in db2.tables["merchant_settlement_day"])
      and db2.tables["merchant_settlement_day"][0]["source_id"] != sid)
s_, d_ = http_error(analyze, db2, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, role="funding")
check("a role that is not a settlement role is refused (funding is a different grain)", s_ == 400 and "role" in str(d_))
s_, d_ = http_error(analyze, db2, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source="some-portal-nobody-configured")
check("a processor that is neither in the portal registry nor on the org's data sources is refused, naming the registry", s_ == 400 and PORTAL in str(d_))
# OVERLAP with the scheduled pull: a day another source already landed would be COUNTED TWICE by the recon → refused
db3 = fresh_db()
db3.seed("merchant_settlement_day", [{"org_id": ORG, "source_id": "99999999-9999-9999-9999-999999999999", "portal_key": PORTAL, "report_key": "card_summary",
                                      "settlement_role": "external_cc", "business_date": "2026-09-01", "merchant_id": MID_A, "card_brand": "visa",
                                      "gross_amount": 1, "refund_amount": 0, "net_amount": 1, "fee_amount": 0, "txn_count": 1, "currency": "USD"}])
m3 = analyze(db3, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, identity=json.dumps(DEC_M))
s_, d_ = http_error(commit, db3, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, identity=json.dumps(DEC_M))
check("NEGATIVE CONTROL: a merchant-day the scheduled pull already landed → the upload is REFUSED (it would double the card recon), the overlap counted",
      m3["verify"]["numbers"]["overlap_with_other_sources"] == 1 and any("ALREADY landed" in r for r in m3["verify"]["refusals"])
      and s_ == 400 and "ALREADY landed" in str(d_) and len(db3.tables["merchant_settlement_day"]) == 1)

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("D. THE CARRIER BILL-PAY REPORT — into the processor feed the coverage recon reads, RE-READ through the mig-939 reader")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
BP_HEADERS = ["Date of Transaction", "Account ID", "Order Number", "User", "Order Type", "Product Name", "Retail Cost", "Merchant Discount"]
BP_ROWS = []
_n = 0
for day in ("2026-09-01", "2026-09-02", "2026-09-03"):
    for acct in ("ACCT-1", "ACCT-2"):
        for i in range(2):
            _n += 1
            BP_ROWS.append([day, acct, f"ORD{_n:04d}", "cashier", "Sales Order", f"RTR Payment ${25 + _n}", f"{25 + _n:.2f}", "0.50"])
        _n += 1
        BP_ROWS.append([day, acct, f"ORD{_n:04d}", "cashier", "Sales Order", "Postpaid Handset", "599.99", "0.00"])   # not a bill payment
        _n += 1
        BP_ROWS.append([day, acct, f"ORD{_n:04d}", "cashier", "Residual", "Monthly residual", "12.00", "0.00"])       # not a bill payment
BP_BILL = money(sum(float(r[6]) for r in BP_ROWS if "RTR" in r[5]))
BP_ALL = money(sum(float(r[6]) for r in BP_ROWS))
BP_CSV = ("\n".join([",".join(BP_HEADERS)] + [",".join(r) for r in BP_ROWS]) + "\n").encode("utf-8")
BP_MAP = {"tx_date": "Date of Transaction", "account_id": "Account ID", "order_number": "Order Number", "user_name": "User",
          "order_type": "Order Type", "product_name": "Product Name", "retail_cost": "Retail Cost", "merchant_discount": "Merchant Discount"}

db = fresh_db()
use(db)
st_ = R.onboarding_intake_state(org_id=ORG)
check("GET /state: the picker (other_kinds), the portal catalog + roles, and the org's bill-pay feed resolution (none configured → the daily-tx layout by default, and it SAYS so)",
      st_["other_kinds"] == list(OI.OTHER_KINDS) and st_["merchant_portals"]["roles"] == list(MP.ROLES)
      and any(c["key"] == PORTAL for c in st_["merchant_portals"]["catalog"])
      and st_["billpay_feed"]["processor"] is None and st_["billpay_feed"]["default_layout"] == "ma_daily_tx" and st_["billpay_feed"]["note"]
      and next(k for k in st_["source_kinds"] if k["value"] == "bill_payments")["layouts"][0]["report_key"] == "ma_daily_tx")
b1 = analyze(db, BP_CSV, "billpay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="ma_daily_tx",
             column_map=json.dumps(BP_MAP), typed_total=str(BP_ALL))
bv = b1["verify"]["numbers"]
check("2.3/2.5: mapped through the daily-tx feed LAYOUT; every row counted; the org's OWN bill-pay predicate (mig 944 order types + product tokens) keeps the RTR rows only, Σ per store-day",
      b1["target_table"] == "raw_ma_daily_tx" and b1["report_key"] == "ma_daily_tx" and bv["rows"] == len(BP_ROWS)
      and bv["billpay_rows"] == 12 and bv["non_billpay_rows"] == 12 and bv["sum_amount"] == BP_BILL and bv["sum_all_rows"] == BP_ALL
      and b1["verify"]["tie"]["match"] is True and bv["processor"] == "vidapay", bv)
st = {r["value"]: r for r in b1["stores"]}
check("2.4: the account ids resolve through the mig-902 map — ACCT-1 → B-100, ACCT-2 unresolved", st["ACCT-1"]["lands_as"] == "B-100" and st["ACCT-2"]["status"] == "unresolved"
      and b1["identity_kind"] == "merchant_id" and b1["merchant_processor"] == "vidapay")
DEC_B = {"store": {"ACCT-2": {"action": "assign", "store_code": "B-200"}}}
c4 = commit(db, BP_CSV, "billpay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="ma_daily_tx",
            column_map=json.dumps(BP_MAP), typed_total=str(BP_ALL), identity=json.dumps(DEC_B))
rmt = db.tables["raw_ma_daily_tx"]
check("COMMIT ok: every row landed through _ingest_mapped_df into raw_ma_daily_tx (period from each row's own date, the slice replaced); the account map row written",
      c4["ok"] is True and c4["saved"] == len(BP_ROWS) and len(rmt) == len(BP_ROWS) and all(r.get("period") for r in rmt)
      and c4["identity_written"]["merchant_ids"] == [("ACCT-2", "B-200")], c4.get("problems"))
cv = c4["verified_numbers"]["numbers"]
check("…RE-READ THROUGH THE MIG-939 READER (_billpay_processor_by_store_day — the coverage recon's own function): the bill-pay Σ equals what was shown; rows = built = inserted",
      cv["feed_read"]["reader"].startswith("_billpay_processor_by_store_day") and cv["sum_amount"] == BP_BILL and cv["rows"] == len(BP_ROWS)
      and c4["verified_numbers"]["rows_landed"] == c4["verified_numbers"]["rows_built"] == c4["verified_numbers"]["rows_inserted"]
      and "mig-939" in c4["verified_numbers"]["basis"], cv)
c4b = commit(db, BP_CSV, "billpay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="ma_daily_tx",
             column_map=json.dumps(BP_MAP), typed_total=str(BP_ALL), identity=json.dumps(DEC_B))
check("re-commit replaces the file's own slice (accounts × dates) — never doubled", c4b["ok"] is True and len(db.tables["raw_ma_daily_tx"]) == len(BP_ROWS))
row_b = next(r for r in db.tables["onboarding_stage_state"] if r["instance_key"].startswith("bill_payments:"))
check("the stage row records the feed table its layout lands in", row_b["payload"]["target_table"] == "raw_ma_daily_tx" and row_b["status"] == "verified")
# a file with no bill-pay row by the org's rule: it would land where the recon reads $0 → refused
BP_NONE = ("\n".join([",".join(BP_HEADERS)] + [",".join(r) for r in BP_ROWS if "RTR" not in r[5]]) + "\n").encode("utf-8")
bn = analyze(db, BP_NONE, "billpay2.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="ma_daily_tx",
             column_map=json.dumps(BP_MAP), typed_total="1", identity=json.dumps(DEC_B))
check("NEGATIVE CONTROL: a report in which the org's rule finds NO bill payment is refused with the reason (it would land, and the recon would read $0.00)",
      bn["verify"]["numbers"]["billpay_rows"] == 0 and any("Not one row" in r for r in bn["verify"]["refusals"]))
# the transaction-detail layout through the feed's own idempotent ingest
EP_HEADERS = ["TransactionID", "TransactionSourceID", "InvoiceID", "SettlementDate", "TerminalID", "UserName", "Product", "ProductTitle", "Type", "Retail", "Discount", "Cost", "Commission"]
EP_ROWS = [["T1001", "1", "I1", "2026-09-01", "TERM-1", "u1", "P1", "RTR $50", "Sale", "50.00", "0", "48.50", "1.50"],
           ["T1002", "1", "I2", "2026-09-01", "TERM-1", "u1", "P9", "SERVICE FEE", "Sale", "1.00", "0", "0.00", "0.00"],
           ["T1003", "1", "I3", "2026-09-02", "TERM-1", "u1", "P1", "RTR $35", "Sale", "35.00", "0", "34.00", "1.00"]]
EP_CSV = ("\n".join([",".join(EP_HEADERS)] + [",".join(r) for r in EP_ROWS]) + "\n").encode("utf-8")
db5 = fresh_db()
db5.seed("store_merchant_id", [{"org_id": ORG, "store_code": "B-100", "processor": "epay", "merchant_id": "TERM-1", "not_required": False}])
e1 = analyze(db5, EP_CSV, "epay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="epay_daily_tx", typed_total="86.00")
ev = e1["verify"]["numbers"]
check("the transaction-detail layout: the feed's fee rule keeps the two payment lines (Σ 85.00), the terminal resolves through the mig-902 map",
      e1["target_table"] == "raw_epay_daily_tx" and ev["billpay_rows"] == 2 and ev["sum_amount"] == 85.0 and ev["sum_all_rows"] == 86.0
      and {r["value"]: r["lands_as"] for r in e1["stores"]} == {"TERM-1": "B-100"} and e1["verify"]["tie"]["match"] is True, ev)
c5 = commit(db5, EP_CSV, "epay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="epay_daily_tx", typed_total="86.00")
check("COMMIT ok through epay_ingest.ingest (idempotent on transaction id), the fee line flagged, store_code stamped; RE-READ through the mig-939 reader ties 85.00",
      c5["ok"] is True and c5["saved"] == 3 and len(db5.tables["raw_epay_daily_tx"]) == 3
      and sum(1 for r in db5.tables["raw_epay_daily_tx"] if r["is_fee"]) == 1 and all(r["store_code"] == "B-100" for r in db5.tables["raw_epay_daily_tx"])
      and c5["verified_numbers"]["numbers"]["sum_amount"] == 85.0 and c5["verified_numbers"]["numbers"]["processor"] == "epay", c5.get("problems"))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("E. BILL PAYMENTS EXTRACTED FROM THE SALES EXPORT — the ONE predicate's lines, exactly; the Bill Payments report")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
SALES_HEADERS = ["Invoice #", "Invoiced At", "Sold By", "Sold On", "Product SKU", "Tracking #", "Product Name", "Category", "Total Price", "Gross Profit", "Refund", "Quantity"]


def sales_lines():
    out = []
    reps = ["Alice Rep", "bob.s"]
    n = 0
    for inv in range(1, 25):
        st = STORE_ADDR if inv % 2 else "B-200"
        for line in range(1, (inv % 3) + 2):
            n += 1
            price = round(19.99 + n * 3.37, 2)
            billpay = (n % 4 == 0)                     # every 4th line is a bill payment under the house rule (department 'rtr')
            refund = (n % 12 == 0)                    # 12, 24, 36, 48 — every one also a bill-pay line, so the void rule is exercised
            if refund:
                price = -price
            out.append({"inv": f"INV-{1000 + inv}", "store": st, "rep": reps[inv % 2], "day": f"08/{1 + inv % 28:02d}/2026 14:{n % 60:02d}:05",
                        "sku": f"SKU-{n % 7}", "imei": f"35{n:013d}", "name": ("Bill Payment $50" if billpay else f"Item {n % 9}"),
                        "cat": ("RTR >> Bill Payments" if billpay else ("Devices >> Phones >> Prepaid" if n % 2 else "Accessories >> Cases")),
                        "price": price, "gp": round(price * 0.31, 2), "refund": "Yes" if refund else "No", "qty": -1 if refund else 1, "billpay": billpay})
    return out


LINES = sales_lines()
SALES_TOTAL = money(sum(l["price"] for l in LINES))
BILL_LINES = [l for l in LINES if l["billpay"] and l["refund"] == "No"]
BILL_SUM = money(sum(l["price"] for l in BILL_LINES))


def sales_xlsx():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Product Sales"
    ws.append(["Product Sales Report"])
    ws.append(SALES_HEADERS)
    for l in LINES:
        ws.append([l["inv"], l["store"], l["rep"], l["day"], l["sku"], l["imei"], l["name"], l["cat"], f"{l['price']:.2f}", f"{l['gp']:.2f}", l["refund"], l["qty"]])
    ws.append(["", "", "", "", "", "", "", "", f"{SALES_TOTAL:.2f}", "", "", ""])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SALES_MAP = {"trans_id": "Invoice #", "store": "Invoiced At", "salesperson": "Sold By", "trans_date": "Sold On", "sku": "Product SKU",
             "serial_1": "Tracking #", "product_desc": "Product Name", "category": "Category", "department": "Category",
             "ext_price": "Total Price", "gp": "Gross Profit", "voided": "Refund", "quantity": "Quantity"}
db = fresh_db()
# PIN UPDATED 2026-09-20 (landing identity): the bill-pay extraction is a LINE-LEVEL reading (department / category /
# product name per sale line), so this export commits under the line-level card — kind `sales`, layout `sales`, which
# now carries the second POS shape's spellings (Invoice # / Sold By / Sold On / Tracking # / Product SKU / Quantity).
# The product layout (`pos_product_sales`) lands in its OWN table and runs no extraction (nothing to match on).
c6 = commit(db, sales_xlsx(), "sales.xlsx", source_kind="sales", pos_source="rq", layout="sales", column_map=json.dumps(SALES_MAP))
bx = c6["verified_numbers"].get("billpay_extract") or {}
check("the sales export commits (ok) and the commit RAN the extraction over the landed slice — the count and Σ of the department-'rtr' lines, non-voided, exactly",
      c6["ok"] is True and bx.get("basis") and bx["lines"] == len(BILL_LINES) and bx["sum"] == BILL_SUM and bx["feed_present"] is False
      and bx["rule_source"] in ("default", "tenant", "carrier_preset"), (c6.get("problems"), bx))
check("…with PROVENANCE: every extracted line matched by the department token 'rtr'; the rule's category tokens matched nothing and are LISTED (the silent-zero signal), not widened",
      bx["per_token"] == [{"by": "department", "token": "rtr", "amount": BILL_SUM, "count": len(BILL_LINES)}]
      and {t["token"] for t in bx["token_coverage"]["unmatched_tokens"]} == {"rtr product", "other carr. payments"}, bx.get("per_token"))
use(db)
landed = [r for r in db.tables["raw_sales"] if r["org_id"] == ORG]
cfg = R._exec_metric_config(db, ORG)
cells = R._sales_cell_agg(landed, R._accessory_config(db, ORG), exec_cfg=cfg)
agg_qty = sum(a.get("bill_qty", 0) for a in cells.values())
agg_amt = money(sum(a.get("bill_amt", 0.0) for a in cells.values()))
check("THE EQUALITY THAT MATTERS: the extraction equals _sales_cell_agg's own bill_qty / bill_amt over the same landed rows — one predicate, not a sibling",
      agg_qty == bx["lines"] and agg_amt == bx["sum"], (agg_qty, agg_amt, bx["lines"], bx["sum"]))
rep = R.billpay_extract_report("August 2026", org_id=ORG)
check("GET /commcalc/billpay-extract/{period}: the derived Bill Payments report — lines with provenance, Σ + count per store-day, the rule and where to edit it, no feed → said so",
      rep["totals"]["lines"] == len(BILL_LINES) and rep["totals"]["sum"] == BILL_SUM and rep["basis"]["feed_present"] is False
      and all(l["matched_by"] == "department" and l["matched_token"] == "rtr" for l in rep["lines"])
      and len(rep["per_store_day"]) == len({(l["store"], l["day"][:5]) for l in BILL_LINES}) and rep["rule"]["edit"].startswith("/commcalc/exec/mtd")
      and all(r["status"] == "sales_only" for r in rep["per_store_day"]), (rep["totals"], rep["basis"]))
check("the Stage-4 row for the sales export carries the extraction note", any("bill-pay line(s) extracted" in (r.get("note") or "")
      for r in R.onboarding_intake_state(org_id=ORG)["rail"]["verify_table"]))
# a voided bill-pay line is not extracted (the same skip rule as the aggregation)
check("a refunded bill-pay line is NOT extracted (the aggregation's own void rule)",
      len([l for l in LINES if l["billpay"] and l["refund"] == "Yes"]) >= 1 and bx["lines"] == len(BILL_LINES))
# now the CARRIER's report is present for some of those days → both figures and the difference
db.seed("store_merchant_id", [])
feed_rows = []
for l in BILL_LINES[:3]:
    mo, dd = l["day"][:2], l["day"][3:5]
    feed_rows.append({"org_id": ORG, "period": "August 2026", "period_month": 8, "period_year": 2026, "tx_date": f"2026-{mo}-{dd}",
                      "account_id": "ACCT-1", "order_number": f"F{dd}", "order_type": "Sales Order", "product_name": "RTR Payment",
                      "retail_cost": l["price"], "merchant_discount": 0})
db.seed("raw_ma_daily_tx", feed_rows)
db.seed("data_source", [{"org_id": ORG, "processor": "vidapay", "label": "carrier processor", "enabled": True}])
rep2 = R.billpay_extract_report("August 2026", org_id=ORG)
check("when the carrier's OWN bill-pay report is present for the period, the report shows BOTH figures and the difference per store-day, and which reader produced the feed side",
      rep2["basis"]["feed_present"] is True and rep2["basis"]["feed_processor"] == "vidapay" and rep2["totals"]["sum_feed"] == money(sum(r["retail_cost"] for r in feed_rows))
      and rep2["totals"]["sum"] == BILL_SUM and any(r["status"] in ("match", "differs", "feed_only") for r in rep2["per_store_day"])
      and rep2["basis"]["feed_reader"].startswith("_billpay_processor_by_store_day"), (rep2["totals"], [r for r in rep2["per_store_day"] if r["status"] != "sales_only"][:3]))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("F. INVENTORY + THE ACTIVATION FEED — activated but still on hand: by IMEI, by mobile THROUGH the sale line, and the unpairable one")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
def imei(tag):
    return (str(tag) + "0" * 15)[:15]


U_SOLD, U_REFUND, U_ACT, U_QUIET, U_TWO = imei("35SOLD1"), imei("35REFUND"), imei("35ACTIV"), imei("35QUIET"), imei("35TWO")
INV_HEADERS = ["Product SKU", "Tracking #", "Product Name", "Location", "Unit Cost", "Total Cost", "Quantity", "Status"]
INV = [(U_SOLD, 300.0), (U_REFUND, 250.0), (U_ACT, 400.0), (U_QUIET, 100.0), (U_TWO, 150.0)]


def inventory_csv():
    lines = [",".join(INV_HEADERS)]
    for k, cost in INV:
        lines.append(f"SKU-1,{k},Phone,{STORE_ADDR},{cost:.2f},{cost:.2f},1,In stock")
    lines.append(f",,,,,{sum(c for _k, c in INV):.2f},{len(INV)},")
    return ("\n".join(lines) + "\n").encode("utf-8")


INV_MAP = {"sku": "Product SKU", "imei": "Tracking #", "item": "Product Name", "store": "Location", "unit_cost": "Unit Cost",
           "total_cost": "Total Cost", "quantity": "Quantity", "status": "Status"}
db = fresh_db()
MOBILE = "5551234567"
db.seed("raw_sales", [
    # U_SOLD: sold, never returned — the classic phantom
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T1", "trans_date": "2026-08-10", "serial_1": U_SOLD, "quantity": 1, "ext_price": 300},
    # U_REFUND: one transaction — a PLAN line carrying the mobile number (no serial), the handset line, then the refund
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T9", "trans_date": "2026-08-12", "serial_1": "", "mdn": f"({MOBILE[:3]}) {MOBILE[3:6]}-{MOBILE[6:]}", "quantity": 1, "ext_price": 40},
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T9", "trans_date": "2026-08-12", "serial_1": U_REFUND, "quantity": 1, "ext_price": 250},
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T10", "trans_date": "2026-08-13", "serial_1": U_REFUND, "quantity": -1, "ext_price": -250},
    # a number that appears on lines for TWO different devices — ambiguous, must be unpairable
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T20", "trans_date": "2026-08-14", "serial_1": U_TWO, "mdn": "5559990000", "quantity": 1, "ext_price": 150},
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T21", "trans_date": "2026-08-15", "serial_1": imei("35OTHER"), "mdn": "5559990000", "quantity": 1, "ext_price": 150},
    {"org_id": ORG, "period": "August 2026", "store": STORE_ADDR, "salesperson": "Alice Rep", "trans_id": "T21", "trans_date": "2026-08-15", "serial_1": U_TWO, "quantity": -1, "ext_price": -150},
])
ACTS = [
    {"Serial#": U_ACT, "Contract Type": "Activation", "Trans ID": "A1", "Trans Date": "8/16/2026", "SP/PO Name": "Plan A"},            # by IMEI — on hand, NO sale line
    {"Serial#": "", "Contract Type": "Activation", "Trans ID": "A2", "Trans Date": "8/12/2026", "SP/PO Name": "Plan A", "MDN": MOBILE},  # by mobile THROUGH T9's handset line (refunded)
    {"Serial#": "", "Contract Type": "Activation", "Trans ID": "A3", "Trans Date": "8/17/2026", "SP/PO Name": "Plan A"},               # neither — unpairable
    {"Serial#": U_SOLD, "Contract Type": "Port", "Trans ID": "A4", "Trans Date": "8/10/2026", "SP/PO Name": "Plan A"},                # NEGATIVE CONTROL: a sales hit — evidence, never a second row
    {"Serial#": "", "Contract Type": "Activation", "Trans ID": "A5", "Trans Date": "8/18/2026", "SP/PO Name": "Plan A", "MDN": "5550001111"},  # a number no sale line carries — unpairable
    {"Serial#": "", "Contract Type": "Activation", "Trans ID": "A6", "Trans Date": "8/18/2026", "SP/PO Name": "Plan A", "MDN": "5559990000"},  # ambiguous number — unpairable
]
db.seed("raw_custom_import", [{"org_id": ORG, "report_key": "activation_details", "period": "August 2026", "source_filename": "acts.xlsx",
                               "row_index": i, "data": d} for i, d in enumerate(ACTS)])
i1 = analyze(db, inventory_csv(), "inventory.csv", source_kind="inventory", pos_source="rq", layout="pos_inventory_listing", column_map=json.dumps(INV_MAP), as_of_date="2026-08-31")
c7 = commit(db, inventory_csv(), "inventory.csv", source_kind="inventory", pos_source="rq", layout="pos_inventory_listing", column_map=json.dumps(INV_MAP), as_of_date="2026-08-31")
sc = c7["verified_numbers"].get("sold_check") or {}
check("the inventory listing commits (ok, 5 units through the snapshot writer) and the commit RAN the auto-check: 2 activated-but-on-hand, 1 sold-not-cleared, 3 unpairable activations",
      c7["ok"] is True and c7["saved"] == 5 and sc.get("basis") and sc["activated_not_rung_out"] == 2 and sc["activated_by_mobile"] == 1
      and sc["sold_not_cleared"] == 1 and sc["activations_unpairable"] == 3 and sc["activations_present"] is True and sc["activations_considered"] == 6
      and sc["activations_carry_device_key"] is True and sc["activations_carry_mobile"] is True, (c7.get("problems"), sc))
use(db)
vt = R.onboarding_intake_state(org_id=ORG)["rail"]["verify_table"]
check("the Stage-4 verify row for inventory says it: '2 unit(s) activated but still on hand'",
      any("2 unit(s) activated but still on hand" in (r.get("note") or "") and "3 activation(s) could not be paired" in r["note"] for r in vt),
      [r.get("note") for r in vt])
out = R.inventory_sold_recon_endpoint(org_id=ORG)
by = {(r["finding"], r["device_key"]): r for r in out["rows"]}
check("GET /inventory-sold-recon: U_ACT is ACTIVATED_NOT_RUNG_OUT by device key (the activation line cited); U_REFUND by mobile number THROUGH the sale line, sale_state sold_then_refunded",
      by[(ISR.ACTIVATED_NOT_RUNG_OUT, U_ACT)]["pairing"] == ISR.PAIR_DEVICE_KEY and by[(ISR.ACTIVATED_NOT_RUNG_OUT, U_ACT)]["evidence"][0]["line"] == "A1"
      and by[(ISR.ACTIVATED_NOT_RUNG_OUT, U_ACT)]["sale_state"] == "no_sale_line"
      and by[(ISR.ACTIVATED_NOT_RUNG_OUT, U_REFUND)]["pairing"] == ISR.PAIR_MOBILE_VIA_SALES and by[(ISR.ACTIVATED_NOT_RUNG_OUT, U_REFUND)]["sale_state"] == "sold_then_refunded"
      and any(e["source"] == "activation" and e["line"] == "A2" and e["via"] == "same transaction" for e in by[(ISR.ACTIVATED_NOT_RUNG_OUT, U_REFUND)]["evidence"]),
      out["rows"])
check("NEGATIVE CONTROL — no double count: U_SOLD is ONE row (sold_not_cleared) with the activation as EVIDENCE, and no activated_not_rung_out row exists for it",
      (ISR.SOLD_NOT_CLEARED, U_SOLD) in by and (ISR.ACTIVATED_NOT_RUNG_OUT, U_SOLD) not in by and by[(ISR.SOLD_NOT_CLEARED, U_SOLD)]["also_activated"] is True
      and any(e["source"] == "activation" and e["line"] == "A4" for e in by[(ISR.SOLD_NOT_CLEARED, U_SOLD)]["evidence"])
      and out["totals"]["to_clear"] == 1 and out["totals"]["activated_not_rung_out"] == 2)
unp = {u["line"]: u for u in out["activations"]["unpairable"]}
check("the unpairable activations are REPORTED with their reason, never guessed: no key & no number; a number no sale line carries; an ambiguous number (candidates listed)",
      unp["A3"]["reason"] == ISR.UNPAIR_NO_KEY and unp["A5"]["reason"] == ISR.UNPAIR_NO_SALE and unp["A6"]["reason"] == ISR.UNPAIR_AMBIGUOUS
      and set(unp["A6"]["candidates"]) == {U_TWO, imei("35OTHER")} and out["totals"]["activations_unpairable"] == 3)
check("the quiet unit (never sold, never activated) and U_TWO (sold, refunded, its number ambiguous) are NOT reported",
      not any(k == U_QUIET or k == U_TWO for _f, k in by))
check("the endpoint states its BASIS per source — what each side is and what it pairs on; activations read ok",
      out["basis"]["sources"]["activations"]["carries_device_key"] is True and out["basis"]["sources"]["activations"]["read_ok"] is True
      and "mobile" in out["basis"]["sources"]["activations"]["pairs_on"] and out["basis"]["sources"]["sales"]["rows"] == 7 and out["activations"]["rows"] == 6)
# no activation feed at all: the report SAYS so and pairs on sales only
db_na = fresh_db()
db_na.tables["raw_custom_import"] = None
use(db_na)
out_na = R.inventory_sold_recon_endpoint(org_id=ORG)
check("without an activation feed the report says 'not present' (never zero activations) and the sales check still runs",
      out_na["activations"]["present"] is False and out_na["basis"]["activations_read_ok"] is False and out_na["totals"]["activated_not_rung_out"] == 0)
# the pure module: the activation line for a sales hit adds no row (pinned again, pure)
_pure = ISR.reconcile([{"serial_1": U_SOLD, "quantity": 1, "trans_id": "T1"}], [{"imei": U_SOLD, "status": "In Stock", "total_cost": 1}], DKEY,
                      activation_rows=[{"serial": U_SOLD, "trans_id": "A4"}])
check("PURE: an activation on a sales hit → one finding, evidence 2 (the sale line and the activation line)",
      len(_pure["rows"]) == 1 and _pure["rows"][0]["finding"] == ISR.SOLD_NOT_CLEARED and len(_pure["rows"][0]["evidence"]) == 2)
check("the Stage-B inventory landing IS what the inventory module reads: the endpoint's on-hand side is inventory_aging_device (mig 216 / 294), the snapshot writer's table",
      "inventory_aging_device" in inspect.getsource(R.inventory_sold_recon_endpoint) and "write_inventory_devices(" in inspect.getsource(R._intake_land)
      and out["basis"]["sources"]["inventory"]["table"] == "inventory_aging_device")

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("G. STAGE 4 / STAGE 5 — one row per kind with its note; the runbook names the real destinations")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
rows = [{"stage": "2", "instance_key": "x_report:rq:x_report", "step": "2.6", "status": "verified", "payload": {"source_ref": "rq", "period": "2026-09-01", "target_table": "pos_tender_summary"},
         "verified_numbers": {"tie": {"our_total": 1, "file_total": 1, "difference": 0, "match": True}, "rows_landed": 8}},
        {"stage": "2", "instance_key": f"merchant_payments:{PORTAL}:settlement", "step": "2.6", "status": "verified", "payload": {"source_ref": PORTAL, "target_table": "merchant_settlement_day"},
         "verified_numbers": {"tie": {"our_total": 2, "file_total": 2, "difference": 0, "match": True}, "rows_landed": 5}},
        {"stage": "2", "instance_key": "bill_payments:carrier:ma_daily_tx", "step": "2.6", "status": "verified", "payload": {"source_ref": "carrier", "layout": "ma_daily_tx", "target_table": "raw_ma_daily_tx"},
         "verified_numbers": {"tie": {"our_total": 3, "file_total": 3, "difference": 0, "match": True}, "rows_landed": 24}},
        {"stage": "2", "instance_key": "inventory:rq:pos_inventory_listing", "step": "2.6", "status": "verified", "payload": {"source_ref": "rq"},
         "verified_numbers": {"tie": {"our_total": 4, "file_total": 4, "difference": 0, "match": True}, "sold_check": {"basis": "x", "activated_not_rung_out": 2, "sold_not_cleared": 1, "activations_present": True}}},
        {"stage": "2", "instance_key": "other:file:loyalty_report", "step": "2.6", "status": "needs_input", "payload": {"name": "loyalty report"},
         "verified_numbers": {"basis": "received"}, "blocking_reason": "no destination for 'loyalty report' yet"}]
rail = OI.rail(rows)
vt = {r["instance_key"]: r for r in rail["verify_table"]}
check("Stage-4: a row per kind, green when verified, the inventory row carrying its note; free-text other stays RED with 'no destination'",
      len(vt) == 5 and all(not vt[k]["red"] for k in vt if not k.startswith("other:")) and vt["other:file:loyalty_report"]["red"]
      and "no destination" in vt["other:file:loyalty_report"]["blocking_reason"]
      and vt["inventory:rq:pos_inventory_listing"]["note"] == "2 unit(s) activated but still on hand; 1 sold but still on hand"
      and vt["x_report:rq:x_report"]["label"].startswith("X-report") and vt[f"merchant_payments:{PORTAL}:settlement"]["label"].startswith("Merchant settlement"))
rb = {m["instance_key"]: m for m in rail["runbook"]["monthly"]}
check("Stage-5: the runbook names where each monthly file lands — pos_tender_summary / merchant_settlement_day / the feed table the layout targets — and free text 'no destination'",
      rb["x_report:rq:x_report"]["lands_in"] == "pos_tender_summary" and rb[f"merchant_payments:{PORTAL}:settlement"]["lands_in"] == "merchant_settlement_day"
      and rb["bill_payments:carrier:ma_daily_tx"]["lands_in"] == "raw_ma_daily_tx" and rb["other:file:loyalty_report"]["lands_in"].startswith("(no destination")
      and rb["x_report:rq:x_report"]["mapping_saved"] is True)
check("…and lists the two cross-check reports beside the two links (the links are unchanged)",
      [l["label"] for l in rail["runbook"]["links"]] == ["Sales report", "Commissions"]
      and {r["screen"] for r in rail["runbook"]["reports"]} == {"bill_payments", "inventory_sold_recon"})

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("H. NEGATIVE CONTROLS — the save guarantee per kind, the gates")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
db = fresh_db()
db.drop_inserts["pos_tender_summary"] = True
cx_ = commit(db, XR_FILE, XR_NAME, source_kind="x_report", pos_source="rq", identity=json.dumps(DEC_X))
check("X-report: a landing that silently drops every row → ok:false 'rows landed 0 ≠ rows built 8', the stage row needs_input — never 'Saved 0 rows' green",
      cx_["ok"] is False and any("rows landed 0" in p for p in cx_["problems"]) and cx_["verified_numbers"]["rows_landed"] == 0
      and next(r for r in db.tables["onboarding_stage_state"] if r["instance_key"].startswith("x_report:"))["status"] == "needs_input")
db = fresh_db()
XR_NODATE = "xreport.xlsx"
xn = analyze(db, XR_FILE, XR_NODATE, source_kind="x_report", pos_source="rq", identity=json.dumps(DEC_X))
s_, d_ = http_error(commit, db, XR_FILE, XR_NODATE, source_kind="x_report", pos_source="rq", identity=json.dumps(DEC_X))
check("X-report with NO date in the file name: the close date is None (never today), the commit is refused until one is typed",
      xn["verify"]["numbers"]["close_date"] is None and any("close date" in r for r in xn["verify"]["refusals"]) and s_ == 400 and "close date" in str(d_))
xd = commit(db, XR_FILE, XR_NODATE, source_kind="x_report", pos_source="rq", identity=json.dumps(DEC_X), as_of_date="2026-09-01")
check("…typed close date → lands under it, date_source 'typed'", xd["ok"] is True and xd["verified_numbers"]["numbers"]["date_source"] == "typed"
      and all(r["close_date"] == "2026-09-01" for r in db.tables["pos_tender_summary"]))
s_, d_ = http_error(analyze, db, XR_FILE, "X-Report_09012026-09022026.xlsx", source_kind="x_report", pos_source="rq")
check("a multi-day X-report is refused (one day's drawer)", s_ == 400 and "SINGLE day" in str(d_))
XR_UNK = xr_workbook([(STORE_ADDR, xr_sheet(STORE_ADDR, XR_A, extra_label="Loyalty Points"))])
xu = analyze(db, XR_UNK, XR_NAME, source_kind="x_report", pos_source="rq")
check("an unrecognised tender label is SKIPPED by the parser and REPORTED as a refusal (its money would be missing from every closing recon) — attestable, never silent",
      "Loyalty Points" in (xu["verify"]["numbers"]["parser"]["unmatched_labels"] or []) and any("Loyalty Points" in r for r in xu["verify"]["refusals"])
      and xu["verify"]["tie"]["match"] is False)
s_, d_ = http_error(analyze, db, XR_FILE, XR_NAME, source_kind="x_report", pos_source="")
check("an X-report needs its POS source", s_ == 400 and "pos_source" in str(d_))
db = fresh_db()
db.drop_inserts["merchant_settlement_day"] = True
cm_ = commit(db, MER_CSV, "settlement.csv", source_kind="merchant_payments", pos_source=PORTAL, identity=json.dumps(DEC_M))
check("merchant settlement: a landing that drops rows → ok:false with the row counts, the stage row needs_input — never 'verified'",
      cm_["ok"] is False and any("rows landed" in p for p in cm_["problems"])
      and next(r for r in db.tables["onboarding_stage_state"] if r["instance_key"].startswith("merchant_payments:"))["status"] == "needs_input")
db = fresh_db()
db.drop_every_other["raw_ma_daily_tx"] = True
cb_ = commit(db, BP_CSV, "billpay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="ma_daily_tx",
             column_map=json.dumps(BP_MAP), typed_total=str(BP_ALL), identity=json.dumps(DEC_B))
check("carrier bill-pay: a landing that drops rows → ok:false (rows re-read ≠ built), never verified", cb_["ok"] is False and any("rows landed" in p for p in cb_["problems"]))
db = fresh_db()
s_, d_ = http_error(analyze, db, BP_CSV, "billpay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="pos_product_sales")
check("a bill-pay report may only map through a processor-feed layout", s_ == 400 and "ma_daily_tx" in str(d_))
db.seed("data_source", [{"org_id": ORG, "processor": "epay", "label": "carrier processor", "enabled": True}])
bm = analyze(db, BP_CSV, "billpay.csv", source_kind="bill_payments", pos_source="carrier-processor", layout="ma_daily_tx",
             column_map=json.dumps(BP_MAP), typed_total=str(BP_ALL), identity=json.dumps(DEC_B))
check("a layout that is NOT the feed this org's coverage recon reads → refused with the feed named (it would land where no recon reads it)",
      bm["verify"]["numbers"]["feed"]["processor"] == "epay" and any("coverage recon reads" in r for r in bm["verify"]["refusals"]))
s_, d_ = http_error(analyze, db, b"Loyalty\nDate,Points\n2026-09-01,5\n", "loyalty.csv", source_kind="merchant_payments", pos_source=PORTAL)
check("a file with no settlement columns is refused, not landed as $0", s_ == 400)
# THE OWNER'S BLOCKER (2026-09-20): "Could not read file: There is no item named 'xl/sharedStrings.xml' in the
# archive" — a workbook that DECLARES a string table it never wrote (inline-string exporters). Through the
# intake's reader it now reads, and the repair is SAID on the payload; an HTML table named .xlsx is refused
# with a sentence. (The module and its own harness — xlsx_tolerant.py / harness_xlsx_tolerant.py — prove the
# case-variant and deleted-table shapes; this pins the WIRING into _read_upload_grids.)
def declared_but_absent_sst(contents):
    import zipfile as _zf
    zin = _zf.ZipFile(io.BytesIO(contents)); out = io.BytesIO()
    with _zf.ZipFile(out, "w", _zf.ZIP_DEFLATED) as zout:
        for n in zin.namelist():
            d = zin.read(n)
            if n == "[Content_Types].xml":
                d = d.decode().replace("</Types>", '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>').encode()
            elif n == "xl/_rels/workbook.xml.rels":
                d = d.decode().replace("</Relationships>", '<Relationship Id="rIdSst" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>').encode()
            zout.writestr(n, d)
    return out.getvalue()
import pandas as _pd
_broken = declared_but_absent_sst(sales_xlsx())
try:
    _pd.read_excel(io.BytesIO(_broken), dtype=str, sheet_name=None, header=None); _repro = False
except Exception as _e:
    _repro = "sharedStrings.xml" in str(_e)
db = fresh_db()
ax = analyze(db, _broken, "sales_by_product.xlsx", source_kind="pos", pos_source="rq", layout="pos_product_sales")
check("THE OWNER'S BLOCKER: the stock reader refuses a workbook whose declared string table is absent; the intake READS it (repair 'added_empty', SAID on the payload) with every row",
      _repro and ax["detect"]["xlsx_repair"] == {"filename": "sales_by_product.xlsx", "action": "added_empty"} and ax["detect"]["data_rows"] == len(LINES) + 1
      and ax["verify"]["numbers"]["sum_amount"] == SALES_TOTAL, ax["detect"].get("xlsx_repair"))
s_, d_ = http_error(analyze, db, b"<html><table><tr><td>x</td></tr></table></html>", "export.xlsx", source_kind="pos", pos_source="rq", layout="pos_product_sales")
check("an HTML table saved with an .xlsx name is refused with a sentence (what it is, what to do), not a zip stack trace",
      s_ == 400 and "its content is html" in str(d_) and "save it as .xlsx" in str(d_), d_)
ok_ = analyze(db, sales_xlsx(), "sales.xlsx", source_kind="pos", pos_source="rq", layout="pos_product_sales")
check("a workbook that reads first time is untouched: no repair noted", ok_["detect"]["xlsx_repair"] is None)
o1 = commit(db, b"Loyalty\nDate,Store,Points\n2026-09-01,100 Main St,5\n", "loyalty.csv", source_kind="other", pos_source="", name="Loyalty report")
check("free-text 'other' is UNCHANGED: recorded + red, nothing written into any table", o1["ok"] is False and o1["recorded"] is True
      and not any(db.tables.get(t) for t in ("pos_tender_summary", "merchant_settlement_day", "raw_ma_daily_tx", "raw_epay_daily_tx")))

# ══════════════════════════════════════════════════════════════════════════════════════════════════
section("I. WIRED, REGISTERED, RULE TWO, THE PAGES")
# ══════════════════════════════════════════════════════════════════════════════════════════════════
paths = [rt.path for rt in R.router.routes]
check("the endpoints are mounted: the Bill Payments report, the inventory recon (unchanged path), the intake",
      "/commcalc/billpay-extract/{period}" in paths and "/commcalc/inventory-sold-recon" in paths and "/commcalc/onboarding/intake/commit" in paths)
land_src = inspect.getsource(R._intake_land)
check("the landings are the EXISTING writers — _xreport_land_rows, merchant_portal_sweep.store_settlement, _ingest_mapped_df, epay_ingest.ingest, write_inventory_devices — and no .insert/.upsert into a destination of its own",
      all(t in land_src for t in ("_xreport_land_rows(", "store_settlement(", "_ingest_mapped_df(", "_epi.ingest(", "write_inventory_devices("))
      and not re.search(r"table\(['\"](pos_tender_summary|merchant_settlement_day|raw_ma_daily_tx|raw_epay_daily_tx|raw_sales|inventory_aging_device)['\"]\)\.(insert|upsert)", land_src))
rsrc = inspect.getsource(R)
check("the X-report upload handler and the intake share ONE upsert loop (_xreport_land_rows); the flat path is the only other pos_tender_summary upsert",
      rsrc.count("table('pos_tender_summary').upsert(") == 2 and "_lx = _xreport_land_rows(" in inspect.getsource(R._upload_file_impl))
check("the bill-pay re-read is the mig-939 reader; the extraction rides exec_metric_defs.line_match; the inventory check rides _cr_resolve_activation_details",
      "_billpay_processor_by_store_day(" in inspect.getsource(R._intake_reread_billpay) and "_emd.line_match" in inspect.getsource(R._intake_billpay_extract_after_sales)
      and "_cr_resolve_activation_details(" in inspect.getsource(R._intake_activation_rows) and "_emd.line_match" in inspect.getsource(R.billpay_extract_report))
check("merchant ids resolve and are written through storeops.merchant_ids (mig 902) — no second map; the closing recon's resolver is asserted above",
      "merchant_ids" in inspect.getsource(R._intake_merchant_resolver) and "_mids.upsert(" in inspect.getsource(R._intake_apply_identity))
vocab = re.compile(r"\b(boost|verizon|cricket|metro|vidapay|total\s+wireless|luxelink|novawave|b2bsoft|iqmetrix|rtpos|payanywhere|transfirst|businesstrack|fiserv|tsys)\b", re.I)
srcs = inspect.getsource(OI) + inspect.getsource(BPX) + inspect.getsource(ISR) + "".join(inspect.getsource(f) for f in (
    R._intake_prepare_stage2, R._intake_prepare_xreport, R._intake_prepare_merchant, R._intake_merchant_table, R._intake_merchant_resolver,
    R._intake_commit_stage2, R._intake_land, R._intake_apply_identity, R._intake_reread_xreport, R._intake_reread_merchant,
    R._intake_reread_billpay, R._intake_billpay_extract_after_sales, R._intake_sold_check_after_inventory, R._intake_activation_rows,
    R._intake_merchant_portal_options, R.billpay_extract_report, R.inventory_sold_recon_endpoint, R._xreport_land_rows))
# the two processor-feed KEYS the mig-939 reader already dispatches on are the only literals allowed, and only in that map
allowed = srcs.replace('_BILLPAY_FEED_LAYOUT = {"epay": "epay_daily_tx", "vidapay": "ma_daily_tx"}', "").replace('processor == "vidapay"', "").replace('"epay_daily_tx"', "")
check("RULE TWO: no carrier, POS vendor, tenant or processor VENDOR name in the intake / extraction / recon code (the two feed keys live only in the mig-939 layout map)",
      not vocab.search(allowed), (vocab.search(allowed) or [""])[0] if vocab.search(allowed) else "")
check("the pure modules import no pandas / fastapi / supabase",
      not re.search(r"^\s*(import|from)\s+(pandas|fastapi|supabase)", inspect.getsource(OI) + inspect.getsource(BPX) + inspect.getsource(ISR), re.M))
idx = open(os.path.join(ROOT, "docs", "SYSTEM_DATA_FLOW_INDEX.md"), encoding="utf-8").read()
check("registered in the index: §30.8 Stage C, the Bill Payments report endpoint, the extraction module, the harness, the epay_daily_tx layout",
      all(t in idx for t in ("30.8", "/commcalc/billpay-extract/{period}", "billpay_extract.py", "harness_onboarding_intake_c.py", "epay_daily_tx", "activated_not_rung_out")))
s2 = open(FE_STAGE2, encoding="utf-8").read() if os.path.exists(FE_STAGE2) else ""
page = open(FE_PAGE, encoding="utf-8").read() if os.path.exists(FE_PAGE) else ""
bp = open(FE_BILLPAY, encoding="utf-8").read() if os.path.exists(FE_BILLPAY) else ""
inv = open(FE_INV, encoding="utf-8").read() if os.path.exists(FE_INV) else ""
check("2.0 is a PICKER of known kinds + free text (rendered from other_kinds), with the portal / role and feed-layout inputs; the matrix kinds skip the column step",
      all(t in s2 for t in ("other_kinds", "merchant_portals", "billpay_feed", "settlement_role", "fd.append('role'", "a.matrix")))
check("2.5 shows the per-kind numbers (tender classes per store-day, per merchant-day Σ, the bill-pay rows the org's rule keeps) and 2.6 the re-read; Stage 4 shows the note",
      all(t in s2 for t in ("per_store_day", "per_merchant_day", "billpay_rows", "close_date")) and "r.note" in page)
check("the Bill Payments report page reads /commcalc/billpay-extract and shows provenance per line, the unmatched tokens and where to edit them; registered in the nav / report directory",
      "billpay-extract" in bp and "matched_token" in bp and "unmatched_tokens" in bp and "exec/mtd" in bp
      and "/commcalc/bill-payments" in open(os.path.join(ROOT, "frontend", "src", "lib", "rbac.ts"), encoding="utf-8").read())
check("the Inventory vs Sold page renders the second section (activated but still on hand) with each row's evidence and the per-source basis",
      "activated_not_rung_out" in inv and "evidence" in inv and "sources" in inv and "unpairable" in inv)

print(f"\n══ onboarding intake — Stage C: {_pass} passed, {_fail} failed ══")
if _failures:
    print("FAILED:\n  - " + "\n  - ".join(_failures))
sys.exit(1 if _fail else 0)
