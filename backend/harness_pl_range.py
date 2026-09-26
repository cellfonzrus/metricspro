#!/usr/bin/env python3
"""PROOF — the P&L exported over a MONTH RANGE: one column per month plus a Total (owner 2026-09-26: "also need
the p&L report to be exported for multiple months … all these need to be platform wide"). Index §4c.

THE RULE: a month column IS that month's P&L page, to the cent — `GET /account/pl-range` (`router.get_pl_range`)
loops THE single-month read (`router.pl_single_month`, which `GET /account/pl/{period}` itself returns) over the
months THE one enumeration (`_period.month_range`) lists; `pl_range.assemble` only lays them side by side and the
Total is only the sum of the month cells. DB-free: the REAL router, the REAL statement_filter / coa company and
market resolution, over an in-memory client that genuinely filters; the stored snapshots are built by the REAL
`engine._assemble` over the REAL `coa.PL_SPEC`.

  A. get_pl is byte-identical to its pre-2026-09-26 inline body (the refactor moved code, changed nothing), for
     every scope / filter / month, computed and not.
  B. THE EQUALITY — for every scope (consolidated, each company, a store) and every filter (stores, a market, a
     company composed with a market that empties it), every month column == `get_pl(month)` for that scope and
     filter: every line, every drill row, every section subtotal, Gross Profit, Net Operating Income, Net Income —
     exact equality, and COMPLETE (every line of every month is on the grid; nothing on the grid is invented).
  C. THE TOTAL == the sum of the month cells of the same row, in integer cents (0.10 + 0.20 + 0.70 == 1.00, not
     0.9999…); a row blank in every month has no Total.
  D. ABSENCE IS NOT ZERO — a month never computed is a blank column named in `missing_months` and headed
     '(not computed)'; a line a month does not carry is blank that month; blanks are not in the Total.
  E. MULTI-TENANT — another org's snapshots (same months, same scope keys, different money) never reach the grid;
     that org's own range reads its own money.
  F. ORDER — the page's order (Revenue, COGS, Gross Profit, Opex, Net Operating Income, Other, Net Income); a line
     that exists in one month only sits where that month shows it; drill rows under their line.
  G. THE RANGE — THE one enumeration; either spelling in; 24 months pass, 25 / reversed / garbage are refused 400.
  H. ONE MONTH == the single-month P&L (the column and the Total both).
  I. THE EXPORT LAYOUT — `export_sheet` cells are the grid's cells; `notes_sheet` names the missing month and
     carries the months' own notes; the endpoint ships both.
  J. PLATFORM-WIDE — the scheduled-report registry entry `account_pl_range` returns the endpoint's own sheets and
     refuses a bad range as a CONFIG error; the server renderer writes a missing money cell EMPTY (xlsx) / blank
     (pdf text), never $0.00, and a real zero stays 0.

Run: cd backend && python3 harness_pl_range.py        (no DB, no network; needs the app's dependencies)
"""
import asyncio
import copy
import io
import json
import os
import sys
from decimal import Decimal
from types import SimpleNamespace

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from fastapi import HTTPException  # noqa: E402

from app.modules.account import coa, engine, statement_engine as SE, autocompute  # noqa: E402
from app.modules.account import _period as PD, pl_range as PR, router as R  # noqa: E402

P = F = 0


def check(label, ok, detail=""):
    global P, F
    if ok:
        P += 1
        print("  PASS  " + label)
    else:
        F += 1
        print("  FAIL  %s   %s" % (label, str(detail)[:600]))


def section(t):
    print("\n── %s %s" % (t, "─" * max(0, 96 - len(t))))


# ══ THE IN-MEMORY CLIENT — filters for real (eq / in / like / range / limit), read-only ══════════════════
class _Q:
    def __init__(self, db, table):
        self.db, self.table, self.filters = db, table, []
        self._limit = self._range = None

    def select(self, *_a, **_k): return self
    def eq(self, k, v): self.filters.append(("eq", k, v)); return self
    def neq(self, k, v): self.filters.append(("neq", k, v)); return self
    def in_(self, k, v): self.filters.append(("in", k, list(v))); return self
    def like(self, k, v): self.filters.append(("like", k, v)); return self
    def ilike(self, k, v): self.filters.append(("ilike", k, v)); return self
    def gte(self, k, v): self.filters.append(("gte", k, v)); return self
    def lte(self, k, v): self.filters.append(("lte", k, v)); return self
    def gt(self, k, v): self.filters.append(("gt", k, v)); return self
    def lt(self, k, v): self.filters.append(("lt", k, v)); return self
    def is_(self, k, v): self.filters.append(("is", k, v)); return self
    def order(self, *_a, **_k): return self
    def limit(self, n): self._limit = n; return self
    def range(self, lo, hi): self._range = (lo, hi); return self

    def _ok(self, r):
        for op, k, v in self.filters:
            x = r.get(k)
            if op == "eq" and str(x) != str(v): return False
            if op == "neq" and str(x) == str(v): return False
            if op == "in" and str(x) not in {str(i) for i in v}: return False
            if op == "like" and not (x is not None and str(x).startswith(str(v).rstrip("%"))): return False
            if op == "ilike" and not (x is not None and str(v).strip("%").lower() in str(x).lower()): return False
            if op in ("gte", "lte", "gt", "lt") and x is None: return False
            if op == "gte" and not str(x) >= str(v): return False
            if op == "lte" and not str(x) <= str(v): return False
            if op == "gt" and not str(x) > str(v): return False
            if op == "lt" and not str(x) < str(v): return False
            if op == "is" and not (str(v) == "null" and x is None): return False
        return True

    def execute(self):
        self.db.reads.append((self.table, list(self.filters)))
        out = [copy.deepcopy(r) for r in self.db.tables.get(self.table, []) if self._ok(r)]
        if self._range:
            out = out[self._range[0]:self._range[1] + 1]
        if self._limit is not None:
            out = out[:self._limit]
        return SimpleNamespace(data=out, count=len(out))


class _Schema:
    def __init__(self, db): self.db = db
    def table(self, n): return _Q(self.db, n)
    def rpc(self, *_a, **_k): return SimpleNamespace(execute=lambda: SimpleNamespace(data=[]))


class Client:
    def __init__(self, tables):
        self.tables, self.reads = tables, []
    def schema(self, _n): return _Schema(self)
    def table(self, n): return _Q(self, n)
    def rpc(self, *_a, **_k): return SimpleNamespace(execute=lambda: SimpleNamespace(data=[]))


# ══ FIXTURE — invented street addresses and figures; statements built by the REAL engine ════════════════
ORG, OTHER = "org-pl-range-a", "org-pl-range-b"
S1, S2, S3 = "100 Main St", "200 Oak Ave", "300 Pine Rd"
C1, C2 = "co-north", "co-south"
JUN, JUL, AUG, SEP = "June 2026", "July 2026", "August 2026", "September 2026"


def blank_inputs():
    return {k: {"by_store": {}, "company_wide": 0.0, "detail": {}} for k, *_ in coa.PL_SPEC + coa.BS_SPEC}


def put(inp, key, store, amount, detail=None):
    inp[key]["by_store"][store] = round(inp[key]["by_store"].get(store, 0.0) + amount, 2)
    for d, v in (detail or {}).items():
        inp[key]["detail"][d] = round(inp[key]["detail"].get(d, 0.0) + v, 2)


def month_inputs(month, scale=1.0):
    """Money that CHANGES shape month to month: a detail key only in July, a journal line only in July, a
    revenue journal line only in August, cents that do not add in binary floating point."""
    inp = blank_inputs()
    k = {JUN: 1.0, JUL: 1.1, AUG: 0.9}[month] * scale
    put(inp, "device_rev", S1, round(20000.37 * k, 2))
    put(inp, "device_rev", S2, round(9000.11 * k, 2))
    put(inp, "device_rev", S3, round(4000.05 * k, 2))
    put(inp, "accessory_rev", S1, {JUN: 0.10, JUL: 0.20, AUG: 0.70}[month] * scale)
    put(inp, "device_cost", S1, round(14000.13 * k, 2))
    put(inp, "device_cost", S3, round(2500.00 * k, 2))
    put(inp, "device_rebate", S1, round(-1200.00 * k, 2))
    put(inp, "wages", S1, round(4100.00 * k, 2))
    put(inp, "wages", S2, round(2200.00 * k, 2))
    det = {"Rent": round(2000.0 * k, 2), "Utilities": round(933.21 * k, 2)}
    if month == JUL:
        det["Internet"] = round(89.99 * scale, 2)
    put(inp, "store_opex", S1, round(sum(det.values()), 2), det)
    put(inp, "store_opex", S3, round(1500.0 * k, 2))
    inp["mi_income"]["company_wide"] = round(7500.0 * k, 2)
    journal = []
    if month == JUL:
        journal.append({"statement": "pl", "account_type": "opex", "account_line": "Insurance",
                        "amount": round(410.0 * scale, 2)})
    if month == AUG:
        journal.append({"statement": "pl", "account_type": "revenue", "account_line": "Other income",
                        "amount": round(250.0 * scale, 2)})
    return inp, journal


def assemble(inp, journal, stores, cw, key, label):
    st = engine._assemble(inp, list(journal), coa.PL_SPEC, coa.PL_LABEL, SE.PL_SECTIONS, key, stores, cw)
    st["scope_key"], st["scope_label"] = key, label
    return st


def snapshots(org, scale):
    rows = []
    for m in (JUN, JUL, AUG):
        inp, journal = month_inputs(m, scale)
        for key, label, stores, cw, jr in (
                ("consolidated", "Consolidated (all companies)", None, True, journal),
                ("company:" + C1, "North Co", {S1, S2}, True, journal),
                ("company:" + C2, "South Co", {S3}, False, []),
                ("store:" + S1, S1, {S1}, False, []),
                ("store:" + S2, S2, {S2}, False, []),
                ("store:" + S3, S3, {S3}, False, [])):
            st = assemble(inp, jr, stores, cw, key, label)
            st["period"] = m
            rows.append({"org_id": org, "period": m, "statement_type": "pl", "scope_key": key,
                         "scope_label": label, "payload": st, "computed_at": f"2026-09-2{len(rows) % 9}T10:00:00Z",
                         "narrative": None, "model": "deterministic", "crosscheck_ok": True})
    return rows


TABLES = {
    "account_statements": snapshots(ORG, 1.0) + snapshots(OTHER, 1000.0),
    "companies": [{"org_id": ORG, "id": C1, "name": "North Co"}, {"org_id": ORG, "id": C2, "name": "South Co"},
                  {"org_id": OTHER, "id": C1, "name": "North Co"}],
    "store_companies": [{"org_id": ORG, "store_address": S1, "company_id": C1},
                        {"org_id": ORG, "store_address": S2, "company_id": C1},
                        {"org_id": ORG, "store_address": S3, "company_id": C2}],
    "stores": [{"org_id": ORG, "store_code": "N1", "address": S1, "market": "North"},
               {"org_id": ORG, "store_code": "N2", "address": S2, "market": "North"},
               {"org_id": ORG, "store_code": "S1", "address": S3, "market": "South"},
               {"org_id": OTHER, "store_code": "X1", "address": S1, "market": "South"}],
}
db = Client(copy.deepcopy(TABLES))
R.sb = lambda: db


def run(coro):
    return asyncio.run(coro)


def get_pl(month, scope="consolidated", stores="", markets="", org=ORG):
    return run(R.get_pl(month, scope=scope, stores=stores, markets=markets, org_id=org))


def get_range(lo, hi, scope="consolidated", stores="", markets="", org=ORG):
    return run(R.get_pl_range(period_from=lo, period_to=hi, scope=scope, stores=stores, markets=markets, org_id=org))


def old_get_pl(period, scope="consolidated", stores="", markets="", org_id=ORG):
    """The pre-2026-09-26 inline body of router.get_pl, verbatim (the byte-identity oracle)."""
    if (stores or "").strip() or (markets or "").strip():
        return R._filtered_read(period, "pl", scope, stores, markets, org_id)
    row = R._read(period, "pl", scope, org_id)
    stale = autocompute.staleness(R.sb(), org_id, period, computed_at=(row.get("computed_at") if row else None))
    if not row:
        return {"period": period, "scope": scope, "computed": False, **stale}
    return {"period": period, "scope": scope, "computed": True,
            "statement": row["payload"], "narrative": row.get("narrative"),
            "model": row.get("model"), "crosscheck_ok": row.get("crosscheck_ok"), **stale}


VIEWS = [("consolidated", "", ""), ("company:" + C1, "", ""), ("company:" + C2, "", ""), ("store:" + S1, "", ""),
         ("consolidated", S1 + "|" + S3, ""), ("consolidated", "", "North"),
         ("company:" + C1, "", "South"), ("consolidated", "", "south")]
MONTHS = [JUN, JUL, AUG, SEP]


def cents(v):
    return int(Decimal(repr(float(v))).quantize(Decimal("0.01")) * 100)


def line_map(st):
    """{(section, key, occurrence): line} for one month's statement — the independent reading of the page."""
    out = {}
    for s in (st or {}).get("sections") or []:
        seen = {}
        for ln in s.get("lines") or []:
            k = str(ln.get("key") or ln.get("label") or "")
            n = seen.get(k, 0)
            seen[k] = n + 1
            out[(s.get("type"), k, n)] = ln
    return out


def grid_lines(grid):
    """{(section, key, occurrence): row} for the grid's line rows (occurrence counted like the page)."""
    out, seen = {}, {}
    for r in grid["rows"]:
        if r["kind"] != "line":
            continue
        k = (r["section"], r["key"])
        n = seen.get(k, 0)
        seen[k] = n + 1
        out[(r["section"], r["key"], n)] = r
    return out


# ══ A. get_pl is byte-identical to its old inline body ═════════════════════════════════════════════════
section("A. get_pl == its pre-change inline body (the refactor moved code, changed nothing)")
same = True
for scope, stores, markets in VIEWS:
    for m in MONTHS + ["2026-07"]:
        a = json.dumps(get_pl(m, scope, stores, markets), sort_keys=True, default=str)
        b = json.dumps(old_get_pl(m, scope, stores, markets), sort_keys=True, default=str)
        if a != b:
            same = False
            print("     differs:", scope, stores, markets, m)
check("A1 every scope / filter / month (computed, not computed, numeric spelling): JSON-identical", same)
check("A2 the page's handler returns the shared single-month read",
      json.dumps(get_pl(JUL, "company:" + C1), sort_keys=True, default=str)
      == json.dumps(R.pl_single_month(JUL, "company:" + C1, "", "", ORG), sort_keys=True, default=str))

# ══ B + C + D. THE EQUALITY, THE TOTAL, ABSENCE ═════════════════════════════════════════════════════════
section("B/C/D. every month column == that month's P&L; Total == Σ months; absence is blank")
def verify(g, sts):
    """The independent reading: each grid cell against each month's single-month statement. Returns the failed
    facts (empty when the grid IS the months)."""
    bad = set()
    gl = grid_lines(g)
    for i, st in enumerate(sts):
        lm = line_map(st)
        for lid, ln in lm.items():                                    # every line of the month is on the grid …
            row = gl.get(lid)
            if row is None:
                bad.add("complete")
                continue
            if row["amounts"][i] != float(ln.get("amount")):          # … with the month's own amount, exactly
                bad.add("lines")
            ix = g["rows"].index(row)
            drill = {}
            for d in g["rows"][ix + 1:]:
                if d["kind"] != "detail":
                    break
                drill[d["label"]] = d["amounts"][i]
            for dk, dv in (ln.get("detail") or {}).items():
                if drill.get(dk) != float(dv):
                    bad.add("detail")
            for dk, dv in drill.items():
                if dv is not None and dk not in (ln.get("detail") or {}):
                    bad.add("detail")
        for lid, row in gl.items():                                   # … and nothing on the grid is invented
            if (st is None or lid not in lm) and row["amounts"][i] is not None:
                bad.add("lines")
        for hk in ("gross_profit", "net_operating_income", "net_income"):
            hr = next(r for r in g["rows"] if r["kind"] == "total" and r["key"] == hk)
            if hr["amounts"][i] != (None if st is None else float(st[hk])):
                bad.add("heads")
        for s in (st or {}).get("sections") or []:
            sr = next(r for r in g["rows"] if r["kind"] == "subtotal" and r["section"] == s["type"])
            if sr["amounts"][i] != float(s["subtotal"]):
                bad.add("subtotal")
    for r in g["rows"]:
        if r["kind"] == "section":
            continue
        cells = [v for v in r["amounts"] if v is not None]
        if r["total"] != (None if not cells else sum(cents(v) for v in cells) / 100):
            bad.add("total")
    return bad


for scope, stores, markets in VIEWS:
    tag = f"[{scope}{' stores=' + stores if stores else ''}{' markets=' + markets if markets else ''}]"
    g = get_range(JUN, SEP, scope, stores, markets)
    sts = [s.get("statement") if s.get("computed") else None for s in (get_pl(m, scope, stores, markets)
                                                                     for m in MONTHS)]
    bad = verify(g, sts)
    check(f"B1 {tag} every line cell == that month's P&L line, exactly", "lines" not in bad)
    check(f"B2 {tag} every drill row == that month's line detail", "detail" not in bad)
    check(f"B3 {tag} every subtotal == that month's section subtotal", "subtotal" not in bad)
    check(f"B4 {tag} Gross Profit / Net Operating Income / Net Income == that month's", "heads" not in bad)
    check(f"B5 {tag} complete: every line of every month is on the grid", "complete" not in bad)
    check(f"C1 {tag} every row's Total == the sum of its month cells (integer cents)", "total" not in bad)
    check(f"D1 {tag} September (never computed) is a blank column, named missing",
          g["missing_months"] == [SEP] and all(r["amounts"][3] is None for r in g["rows"])
          and g["month_status"][3]["computed"] is False, g["missing_months"])

# the views above are not vacuous: the filters BIND stores and carry real money
gn = get_range(JUN, AUG, "consolidated", "", "North")
check("B6 the market filter binds its two stores every month, with money on the grid",
      [s["matched_stores"] for s in gn["month_status"]] == [2, 2, 2]
      and next(r for r in gn["rows"] if r["key"] == "device_rev")["total"] > 0,
      [s["matched_stores"] for s in gn["month_status"]])
gc = get_range(JUN, AUG, "company:" + C2)
check("B7 a company scope reads its own snapshot (South Co: one store's money, not the consolidated)",
      gc["scope_label"] == "South Co" and next(r for r in gc["rows"] if r["key"] == "device_rev")["amounts"]
      == [4000.05, round(4000.05 * 1.1, 2), round(4000.05 * 0.9, 2)],
      next(r for r in gc["rows"] if r["key"] == "device_rev")["amounts"])

# NEGATIVE CONTROLS — the reading above turns red on a grid that is not the months
g0 = get_range(JUN, SEP)
sts0 = [s["statement"] if s["computed"] else None for s in (get_pl(m) for m in MONTHS)]
check("B8 control: the untampered grid verifies clean", verify(g0, sts0) == set(), verify(g0, sts0))
t1 = copy.deepcopy(g0)
next(r for r in t1["rows"] if r["kind"] == "line" and r["key"] == "wages")["amounts"][1] += 0.01
check("B9 control: one cent off one line in one month -> RED", "lines" in verify(t1, sts0))
t2 = copy.deepcopy(g0)
next(r for r in t2["rows"] if r["kind"] == "total" and r["key"] == "net_income")["total"] += 0.01
check("B10 control: a Total one cent off the sum -> RED", "total" in verify(t2, sts0))
t3 = copy.deepcopy(g0)
t3["rows"] = [r for r in t3["rows"] if r.get("label") != "Insurance"]
check("B11 control: a dropped line -> RED (incomplete)", "complete" in verify(t3, sts0))
t4 = copy.deepcopy(g0)
next(r for r in t4["rows"] if r["kind"] == "line" and r["key"] == "wages")["amounts"][3] = 0.0
check("B12 control: $0.00 in a month never computed -> RED (absence is not zero)", "lines" in verify(t4, sts0))
_orig_single = R.pl_single_month


def _tampered_single(m, *a, **k):
    d = _orig_single(m, *a, **k)
    if d.get("computed"):
        d["statement"]["sections"][0]["lines"][0]["amount"] = 1.0
    return d


R.pl_single_month = _tampered_single
gt = get_range(JUN, AUG)
R.pl_single_month = _orig_single
check("B13 control: the range really reads THROUGH pl_single_month (a patched single read moves the grid)",
      next(r for r in gt["rows"] if r["kind"] == "line")["amounts"] == [1.0, 1.0, 1.0])

g = get_range(JUN, SEP)
acc = next(r for r in g["rows"] if r["kind"] == "line" and r["key"] == "accessory_rev")
check("C2 0.10 + 0.20 + 0.70 totals to exactly 1.00 (not a float 0.9999…)", acc["total"] == 1.0, acc)
check("C3 …while each month keeps its own exact cell", acc["amounts"][:3] == [0.1, 0.2, 0.7], acc["amounts"])
ins = next(r for r in g["rows"] if r["kind"] == "line" and r["label"] == "Insurance")
check("D2 a line only July carries is blank in June and August, and its Total is July's alone",
      ins["amounts"] == [None, 410.0, None, None] and ins["total"] == 410.0, ins)
internet = next(r for r in g["rows"] if r["kind"] == "detail" and r["label"] == "Internet")
check("D3 a drill row only July carries is blank in the other months", internet["amounts"] == [None, 89.99, None, None],
      internet)
check("D4 a row with no number in any month has no Total (never $0.00)",
      PR._sum_cents([None, None]) is None and PR._sum_cents([0.0, None]) == 0.0)

# ══ E. MULTI-TENANT ═══════════════════════════════════════════════════════════════════════════════════
section("E. another org's snapshots never reach the grid")
dev = next(r for r in g["rows"] if r["kind"] == "line" and r["key"] == "device_rev")
ours = [get_pl(m)["statement"] for m in (JUN, JUL, AUG)]
want_dev = [next(l for s in st["sections"] for l in s["lines"] if l["key"] == "device_rev")["amount"] for st in ours]
check("E1 the org's grid carries its own money", dev["amounts"][:3] == want_dev, (dev["amounts"], want_dev))
check("E2 …never the other org's (1000× the money)", all(v < 100000 for v in dev["amounts"][:3]), dev["amounts"])
go = get_range(JUN, AUG, org=OTHER)
dev_o = next(r for r in go["rows"] if r["kind"] == "line" and r["key"] == "device_rev")
check("E3 the other org's range reads its own money", all(v > 1000000 for v in dev_o["amounts"]), dev_o["amounts"])
reads = [f for (t, f) in db.reads if t == "account_statements"]
check("E4 every snapshot read carries an org_id filter", reads and all(any(op == "eq" and k == "org_id"
                                                                           for op, k, _v in f) for f in reads))
gs = get_range(JUN, AUG, "company:" + C1, "", "South")
check("E5 a company composed with a market it has no store in reads NOTHING of either (fail closed)",
      all(v in (None, 0.0) for r in gs["rows"] if r["kind"] == "line" for v in r["amounts"]))

# ══ F. ORDER ════════════════════════════════════════════════════════════════════════════════════════════
section("F. the page's order; a one-month line sits where that month shows it")
seq = [(r["kind"], r["section"] if r["kind"] != "total" else r["key"]) for r in g["rows"]
       if r["kind"] in ("section", "subtotal", "total")]
check("F1 Revenue, COGS, Gross Profit, Opex, Net Operating Income, Other, Net Income", seq == [
    ("section", "revenue"), ("subtotal", "revenue"), ("section", "cogs"), ("subtotal", "cogs"),
    ("total", "gross_profit"), ("section", "opex"), ("subtotal", "opex"), ("total", "net_operating_income"),
    ("section", "other"), ("subtotal", "other"), ("total", "net_income")], seq)
opex_keys = [r["key"] for r in g["rows"] if r["kind"] == "line" and r["section"] == "opex"]
jul_opex = [l["key"] for s in get_pl(JUL)["statement"]["sections"] if s["type"] == "opex" for l in s["lines"]]
check("F2 the opex lines keep July's order (the July-only journal line included)",
      [k for k in opex_keys if k in jul_opex] == jul_opex, (opex_keys, jul_opex))
rev_keys = [r["key"] for r in g["rows"] if r["kind"] == "line" and r["section"] == "revenue"]
aug_rev = [l["key"] for s in get_pl(AUG)["statement"]["sections"] if s["type"] == "revenue" for l in s["lines"]]
check("F3 an August-only revenue line sits where August shows it", [k for k in rev_keys if k in aug_rev] == aug_rev,
      (rev_keys, aug_rev))
so = next(i for i, r in enumerate(g["rows"]) if r["kind"] == "line" and r["key"] == "store_opex")
check("F4 drill rows sit directly under their line",
      [r["label"] for r in g["rows"][so + 1:so + 4]] == ["Rent", "Utilities", "Internet"],
      [r["label"] for r in g["rows"][so + 1:so + 4]])
check("F5 _merge_order inserts a new key after its predecessor and never moves a placed key",
      PR._merge_order(["a", "c"], ["a", "b", "c"]) == ["a", "b", "c"]
      and PR._merge_order(["a", "b", "c"], ["c", "x"]) == ["a", "b", "c", "x"]
      and PR._merge_order(["a", "b"], ["z", "b"]) == ["z", "a", "b"])

# ══ G. THE RANGE ════════════════════════════════════════════════════════════════════════════════════════
section("G. THE one enumeration; the cap; refusals")
check("G1 the months are exactly _period.month_range's", g["months"] == PD.month_range(JUN, SEP), g["months"])
check("G2 either spelling in, canonical out", get_range("2026-06", "2026-08")["months"] == [JUN, JUL, AUG])
check("G3 24 months pass (the cap)", len(get_range("2024-10", "2026-09")["months"]) == PR.MAX_MONTHS == 24)


def refused(lo, hi):
    try:
        get_range(lo, hi)
    except HTTPException as e:
        return e.status_code == 400 and bool(e.detail)
    return False


check("G4 25 months are refused 400 with the reason", refused("2024-09", "2026-09"))
check("G5 a reversed range is refused 400", refused(AUG, JUN))
check("G6 an unparseable month is refused 400 (never guessed)", refused("someday", AUG))
check("G7 a blank to-month is the one month", get_range(JUL, "")["months"] == [JUL])

# ══ H. ONE MONTH == THE SINGLE-MONTH P&L ═══════════════════════════════════════════════════════════════
section("H. a one-month range is the single-month P&L")
g1 = get_range(JUL, JUL, "store:" + S1)
st = get_pl(JUL, "store:" + S1)["statement"]
lm = line_map(st)
check("H1 every line == the single-month line and the Total == that same number",
      all(r["amounts"] == [float(lm[lid]["amount"])] and r["total"] == float(lm[lid]["amount"])
          for lid, r in grid_lines(g1).items()) and len(grid_lines(g1)) == len(lm))
check("H2 net income column == the page's net income",
      next(r for r in g1["rows"] if r["key"] == "net_income")["amounts"] == [st["net_income"]])

# ══ I. THE EXPORT LAYOUT ═══════════════════════════════════════════════════════════════════════════════
section("I. export_sheet / notes_sheet — the grid's own cells")
sh, notes = g["sheets"]
hdr = [c["header"] for c in sh["columns"]]
check("I1 columns: Section, Line, one per month (the missing one says so), Total",
      hdr == ["Section", "Line", JUN, JUL, AUG, SEP + " (not computed)", "Total"], hdr)
check("I2 month and Total columns are money", [c.get("money", False) for c in sh["columns"]]
      == [False, False, True, True, True, True, True])
body = [r for r in g["rows"] if r["kind"] != "section"]
check("I3 one export row per grid row (section titles ride the Section column)", len(sh["rows"]) == len(body))
check("I4 every export cell == the grid cell; Total == the grid Total",
      all(all(er[PR.month_column_key(i)] == gr["amounts"][i] for i in range(4)) and er["total"] == gr["total"]
          for er, gr in zip(sh["rows"], body)))
check("I5 drill rows read '↳ <detail>', subtotals and headline rows are labelled",
      any(r["line"].strip() == "↳ Rent" for r in sh["rows"])
      and any(r["line"].strip() == "Subtotal — Revenue" for r in sh["rows"])
      and any(r["section"] == "Totals" and r["line"] == "Net Income" for r in sh["rows"]))
check("I6 the notes sheet names the month never computed", any(n["months"] == SEP and "Not computed" in n["note"]
                                                               for n in notes["rows"]), notes["rows"][:3])
gf = get_range(JUN, AUG, "consolidated", S1, "")
check("I7 the months' own notes travel (the store-filter note, once, with its months)",
      any("Store/market filter active" in n["note"] and n["months"] == [JUN, JUL, AUG] for n in gf["notes"]))

# ══ J. PLATFORM-WIDE ═══════════════════════════════════════════════════════════════════════════════════
section("J. the scheduled-report registry and the server renderer")
from app.modules.notify import finance_reports as FR, report_registry as RREG, render as RND  # noqa: E402
from app.modules.notify.report_filters import ReportConfigError  # noqa: E402
check("J1 the registry lists account_pl_range", "account_pl_range" in RREG.REPORTS)
built = run(FR._account_pl_range(ORG, {"period_from": JUN, "period_to": SEP, "scope": "consolidated"}))
check("J2 the registry's sheets ARE the endpoint's sheets (one layout)",
      json.dumps(built["sheets"], sort_keys=True) == json.dumps(g["sheets"], sort_keys=True))
check("J3 …and it says which month was not computed", "not computed: " + SEP in built["subtitle"], built["subtitle"])
try:
    run(FR._account_pl_range(ORG, {"period_from": AUG, "period_to": JUN}))
    j4 = False
except ReportConfigError:
    j4 = True
check("J4 a reversed saved range is a CONFIG error on the schedule, not a crash", j4)
try:
    run(FR._account_pl_range(ORG, {"period_from": "2025-01", "period_to": "2025-02"}))
    j5 = False
except ValueError as e:
    j5 = "not computed" in str(e)
check("J5 a range with no computed month says so instead of sending an empty file", j5)
xb = RND.build_xlsx(built)
from openpyxl import load_workbook  # noqa: E402
ws = load_workbook(io.BytesIO(xb))[built["sheets"][0]["name"][:31]]
hdrs = [c.value for c in ws[1]]
sep_col = hdrs.index(SEP + " (not computed)") + 1
jun_col = hdrs.index(JUN) + 1
check("J6 xlsx: a not-computed month's cells are EMPTY, never 0", all(ws.cell(row=r, column=sep_col).value is None
                                                                   for r in range(2, ws.max_row + 1)))
check("J7 xlsx: computed month cells are the grid's numbers",
      [ws.cell(row=r, column=jun_col).value for r in range(2, ws.max_row + 1)]
      == [er[PR.month_column_key(0)] for er in built["sheets"][0]["rows"]])
zero = RND.build_xlsx({"sheets": [{"name": "Z", "columns": [{"header": "A", "key": "a", "money": True}],
                                    "rows": [{"a": 0.0}, {"a": None}, {"a": 12.5}]}]})
zs = load_workbook(io.BytesIO(zero))["Z"]
check("J8 xlsx: a REAL zero stays 0, a missing value is empty, a number is a number",
      [zs.cell(row=r, column=1).value for r in (2, 3, 4)] == [0, None, 12.5])
col = {"header": "A", "key": "a", "money": True}
check("J9 pdf text: missing → blank, zero → $0.00", RND._display(col, {"a": None}) == ""
      and RND._display(col, {"a": 0}) == "$0.00" and RND._display(col, {"a": 1234.5}) == "$1,234.50")

print("\n%d passed, %d failed" % (P, F))
sys.exit(1 if F else 0)
